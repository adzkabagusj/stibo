"""
╔══════════════════════════════════════════════════════════════════╗
║   STIBO INBOUND XML GENERATOR — Asics EAN v1.0                  ║
║   EAN Source File → Stibo STEP XML (Generic + Variants)         ║
╚══════════════════════════════════════════════════════════════════╝

Input file pattern:
    0999-SP-ASICS-EAN Source-SP2027-ID-1.xlsx

Sheet: "Sheet2"

Column layout (Asics EAN Source):
    STYLE/COLOUR  -> Principal Style Code (used in KEY_InboundArticle)
    Size          -> Size value (used directly in KEY_InboundVariant)
    BARCODE       -> Barcode (EAN/UPC)

Article structure:
  One Generic per unique STYLE/COLOUR
  One Variant per row (size used as-is, no MDD lookup)
  Principal Barcode → Variant DC_Barcode DataContainer

KEY_InboundArticle formula:
    BrandCode + STYLE/COLOUR

KEY_InboundVariant formula:
    KEY_InboundArticle + Size
"""

import re
import sys
import logging
import argparse
import os
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

# ======================================================================
# PATHS
# ======================================================================
BASE_DIR = Path(os.environ.get("LAMBDA_TMP_DIR", Path(__file__).resolve().parent))

EAN_DIR     = BASE_DIR / "input" / "ean"
XML_OUT_DIR = BASE_DIR / "output" / "xml"
LOG_DIR     = BASE_DIR / "output" / "logs"

for _d in (EAN_DIR, XML_OUT_DIR, LOG_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# ======================================================================
# BRAND CONSTANTS
# ======================================================================
BRAND_NAME = "Asics"
BRAND_CODE = "ASC"

# ======================================================================
# XML NAMESPACE
# ======================================================================
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"
ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z]+)?="[^"]*"')


# ======================================================================
# LOADER
# ======================================================================

class AsicsEANLoader:
    """Loads Asics EAN source data from an Excel file.

    Header auto-detection: scans the first 15 rows for signals like
    'style/colour', 'barcode', 'size'.
    """

    HEADER_SIGNALS = {"style/colour", "barcode", "size", "style", "colour"}
    TARGET_SHEET   = "Sheet2"

    def __init__(self, path: Path):
        self.path = path
        self.df = pd.DataFrame()
        self.sheet = ""
        self.headers: list[str] = []
        self._load()

    def _load(self):
        log.info("[Asics-EAN] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        log.info("[Asics-EAN] Available sheets: %s", wb.sheetnames)

        # Find TARGET_SHEET (case-insensitive)
        target = None
        for sn in wb.sheetnames:
            if sn.strip().lower() == self.TARGET_SHEET.lower():
                target = sn
                break
        if not target:
            wb.close()
            raise ValueError(
                f"Sheet '{self.TARGET_SHEET}' not found in {self.path.name}. "
                f"Available: {wb.sheetnames}"
            )
        log.info("[Asics-EAN] Using sheet: '%s'", target)
        ws = wb[target]
        rows = list(ws.iter_rows(values_only=True))

        # Find header row
        hdr_idx = 0
        for i in range(min(15, len(rows))):
            if rows[i]:
                rv = {str(v).strip().lower() for v in rows[i] if v is not None}
                if rv & self.HEADER_SIGNALS:
                    hdr_idx = i
                    break

        log.info("[Asics-EAN] Header at row %d (0-indexed)", hdr_idx)
        header = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[hdr_idx])]

        # Handle duplicate column names
        seen_cols: dict = {}
        for ci, col in enumerate(header):
            if col in seen_cols:
                seen_cols[col] += 1
                header[ci] = f"{col}_{seen_cols[col]}"
            else:
                seen_cols[col] = 0

        df = pd.DataFrame(rows[hdr_idx + 1:], columns=header)

        # Filter rows with valid STYLE/COLOUR
        style_col = self._find_column(header, ["STYLE/COLOUR", "Style/Colour", "style/colour", "STYLE", "Style"])
        if style_col:
            df = df[df[style_col].notna() & (~df[style_col].astype(str).str.strip().isin(["", "None", "nan"]))]

        wb.close()
        self.df = df.reset_index(drop=True)
        self.sheet = target
        self.headers = header
        log.info("[Asics-EAN] %d rows loaded", len(self.df))
        if self.headers:
            log.info("[Asics-EAN] Headers: %s", self.headers[:20])

    @staticmethod
    def _find_column(headers: list, candidates: list) -> str | None:
        for c in candidates:
            if c in headers:
                return c
            for h in headers:
                if h.lower() == c.lower():
                    return h
        return None

    def find_column(self, *candidates) -> str | None:
        return self._find_column(self.headers, list(candidates))


# ======================================================================
# HELPERS
# ======================================================================

def _s(v) -> str:
    """Clean string value."""
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "0", "NaT") else s


def _fmt_barcode(val) -> str:
    """Format barcode/EAN/UPC value."""
    if val is None:
        return ""
    try:
        return str(int(float(str(val).strip())))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s in ("None", "nan") else s


def _clean_style(val) -> str:
    """Clean style/colour code."""
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s in ("None", "nan") else s


def _clean_size(val) -> str:
    """Clean size value — used as-is in KEY_InboundVariant."""
    if val is None:
        return ""
    s = str(val).strip()
    # Remove trailing .0 from numeric sizes read as float (e.g. "10.0" → "10")
    if re.match(r'^\d+\.0$', s):
        s = s[:-2]
    return "" if s in ("None", "nan") else s


# ======================================================================
# MAPPER
# ======================================================================

def group_ean_rows(raw_rows: list[dict], brand_code: str,
                   loader: AsicsEANLoader) -> dict[str, dict]:
    """
    Groups EAN rows into generics.

    KEY_InboundArticle = BrandCode + STYLE/COLOUR
    KEY_InboundVariant = KEY_InboundArticle + Size

    Each unique STYLE/COLOUR = One Generic
    Each row = One Variant (size from 'Size', barcode from 'BARCODE')
    """
    result: dict[str, dict] = {}

    # Locate columns
    style_col   = loader.find_column("STYLE/COLOUR", "Style/Colour", "style/colour", "STYLE", "Style")
    size_col    = loader.find_column("Size", "SIZE", "size", "Grid Value", "GRID VALUE")
    barcode_col = loader.find_column("BARCODE", "Barcode", "barcode", "UPC", "EAN", "EAN Code", "EAN/UPC", "GTIN")

    if not style_col:
        log.error("[Asics-EAN] Cannot find STYLE/COLOUR column")
        return result

    log.info("[Asics-EAN] All available headers: %s", loader.headers)
    log.info("[Asics-EAN] Using columns — STYLE/COLOUR: '%s'  Size: '%s'  BARCODE: '%s'",
             style_col, size_col, barcode_col)

    if not size_col:
        log.warning("[Asics-EAN] Size column NOT FOUND — size will default to empty!")
    if not barcode_col:
        log.warning("[Asics-EAN] BARCODE column NOT FOUND — barcodes will be empty!")

    for row in raw_rows:
        style_raw   = _clean_style(row.get(style_col))
        if not style_raw:
            continue

        size_raw    = _clean_size(row.get(size_col)) if size_col else ""
        barcode_str = _fmt_barcode(row.get(barcode_col)) if barcode_col else ""

        # KEY_InboundArticle = BrandCode + STYLE/COLOUR
        generic_code = f"{brand_code}{style_raw}"

        if generic_code not in result:
            result[generic_code] = {
                "style_code":   style_raw,
                "generic_code": generic_code,
                "brand_code":   brand_code,
                "variants":     {},
            }

        # KEY_InboundVariant = KEY_InboundArticle + Size
        variant_code = f"{generic_code}{size_raw}"

        g = result[generic_code]
        # Use size as key; keep first occurrence if duplicates exist
        if size_raw not in g["variants"]:
            g["variants"][size_raw] = {
                "size":         size_raw,
                "barcode":      barcode_str,
                "variant_code": variant_code,
            }

    log.info("[Mapper] %d rows → %d generics  %d variants",
             len(raw_rows), len(result),
             sum(len(g["variants"]) for g in result.values()))
    return result


# ======================================================================
# XML HELPERS
# ======================================================================

def _add_barcode_datacontainer(parent_el: ET.Element, barcode: str) -> None:
    """Add DC_Barcode DataContainer with barcode attributes."""
    if not barcode:
        return
    dc_root = ET.SubElement(parent_el, f"{{{STIBO_NS}}}DataContainers")
    mdc = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    # AT_Barcode — the actual barcode value
    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = barcode

    # AT_BarcodeType — P for Principal
    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    # AT_MainEANIndicator — Y for Yes
    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


# ======================================================================
# PRODUCT XML BUILDER
# ======================================================================

def build_product_xml(generic: dict) -> str:
    """Build XML for one Generic article with its Variants."""
    if not generic.get("style_code"):
        return ""

    # Generic Article element
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    # KEY_InboundArticle
    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = generic["generic_code"]

    # Empty Values element for generic
    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")

    # Variant Articles — sorted by size for stable output
    for size_key in sorted(generic["variants"].keys()):
        variant = generic["variants"][size_key]

        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        # KEY_InboundVariant
        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant["variant_code"]

        # Empty Values element for variant
        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")

        # Add barcode DataContainer if barcode exists
        if variant.get("barcode"):
            _add_barcode_datacontainer(v_el, variant["barcode"])

    return ET.tostring(g_el, encoding="unicode")


# ======================================================================
# ORCHESTRATOR
# ======================================================================

def run(args, auditor=None):
    """Main entry point for Asics EAN processing."""
    brand      = getattr(args, "brand",      BRAND_NAME)
    brand_code = getattr(args, "brand_code", BRAND_CODE)
    season     = getattr(args, "season",     "SS27")

    log.info("[Asics-EAN] Starting: brand=%s code=%s season=%s", brand, brand_code, season)

    # Load input file
    ean_file = next(EAN_DIR.glob("*.xls*"), None)
    if not ean_file:
        raise FileNotFoundError(f"No EAN source file found in {EAN_DIR}")

    loader = AsicsEANLoader(ean_file)

    if loader.df.empty:
        log.warning("[Asics-EAN] No data rows found")
        return

    raw_rows = loader.df.to_dict("records")
    generics = group_ean_rows(raw_rows, brand_code, loader)

    if not generics:
        log.warning("[Asics-EAN] No valid generics produced")
        return

    log.info("[Asics-EAN] Total generics to write: %d", len(generics))

    # Generate XML
    xml_filename = f"{ean_file.stem}.xml"
    xml_path     = XML_OUT_DIR / xml_filename
    export_time  = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = variant_count = 0

    log.info("[Asics-EAN] Writing XML → %s", xml_filename)
    with open(xml_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n\n'
        )
        f.write("  <Products>\n")

        for generic in generics.values():
            px = build_product_xml(generic)
            if px:
                f.write(f"    {_XMLNS_RE.sub('', px)}\n")
                generic_count += 1
                variant_count += len(generic["variants"])

        f.write("  </Products>\n</STEP-ProductInformation>\n")

    file_kb = xml_path.stat().st_size // 1024
    log.info("[Asics-EAN] XML written: %s  (%d KB)", xml_path.name, file_kb)
    print(
        f"=== ASICS EAN SUMMARY ===\n"
        f"  Input rows : {len(raw_rows)}\n"
        f"  Total generics found : {len(generics)}\n"
        f"  Generics written     : {generic_count}\n"
        f"  Variants written     : {variant_count}\n"
        f"  XML size   : {file_kb} KB",
        flush=True,
    )

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return xml_path


# ======================================================================
# MAIN
# ======================================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Asics EAN Source → Stibo XML")
    parser.add_argument("--brand",      default=BRAND_NAME)
    parser.add_argument("--brand_code", default=BRAND_CODE)
    parser.add_argument("--season",     default="SS27")
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )
    run(args)
