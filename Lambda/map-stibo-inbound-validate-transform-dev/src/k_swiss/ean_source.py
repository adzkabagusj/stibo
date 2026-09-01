"""
╔══════════════════════════════════════════════════════════════════╗
║   STIBO INBOUND XML GENERATOR — K-Swiss Inline EAN  v1.0         ║
║   EAN/UPC source (e.g. "FW26 UPC - PT.MAP 20260508.xlsx")        ║
║   → Stibo STEP XML                                                ║
║   (Structure mirrors dr_marten/inline_ean_source.py)              ║
╚══════════════════════════════════════════════════════════════════╝

Purpose:
    Generates Variant-level STEP XML carrying barcode (EAN/UPC) data
    ONLY. Companion to k_swiss/global_line_order_form.py, which sends
    the Generic-level article data (no barcodes) from the Performance
    Linelist / Lifestyle Linelist sheets.

Column mapping — per the "EAN File Inline--FW26 UPC" Integrations Team
Query and MAA's 12/08/2026 comment on it:
    Principal Style + Color  -> MATERIAL     (matches the Global Line
                                 Order Form's "MATERIAL #" column, e.g.
                                 "04437-031-M" = Style-Color-Width, the
                                 pipe-key barcodes must attach to)
    Size                     -> Europe Size  (one size per row — unlike
                                 the Generic-level "Size Range" span,
                                 e.g. "6H-12, 13, 14")
    EAN                      -> International Article Number (EAN/UPC)

KEY_InboundArticle formula (MUST match global_line_order_form.py's
generic_code exactly, so barcodes attach to the correct Generic
article):
    brand_code[:3].upper() + MATERIAL#          e.g. "KSW04437-031-M"

KEY_InboundVariant formula (same 3-char MAA size-code convention used
across implus/dr_marten EAN sources):
    KEY_InboundArticle + size_code(Size)        e.g. "KSW04437-031-M009"

Input file: K-Swiss EAN/UPC source Excel (e.g. "FW26 UPC - PT.MAP
20260508.xlsx"). Confirmed real column layout (Sheet1):
    International Article Number (EAN/UPC) | Generic Number |
    Generic Description | Europe Size | Material Number | EAN category |
    Division

    IMPORTANT — "Material Number" in this file is NOT the Generic-level
    key despite the name: it's a per-row/per-size code (e.g.
    "04413-134-M-10", one distinct value per of 418 rows). The real
    Generic-level key — the one that matches global_line_order_form.py's
    "MATERIAL #" column verbatim (e.g. "04413-134-M", only 32 distinct
    values across the file) — is the "Generic Number" column. Column
    matching below therefore checks Generic Number FIRST, and only
    falls back to a literal "MATERIAL #" column or a Style+Color(+Width)
    reconstruction if Generic Number isn't present (e.g. a future drop
    renames it):
    Generic  : "Generic Number" / "Generic No" / "Generic #" / "Generic Code"
               (checked first — this is the true Generic/Article key)
    Style    : "Principal Style Code" / "Principal Style" / "Style Code" / "Style"
    Color    : "Principal Color Code" / "Color Code" / "Colour Code" / "Color"
    Width    : "Width" (optional — MATERIAL # ends in -M/-W)
    Material : "MATERIAL #" / "Material #"  (literal-hash form only —
               "Material Number"/"Material No" are deliberately excluded
               from this alias list; in this brand's file that phrasing
               is the per-variant code, not the Generic key)
    Size     : "Size" / "Europe Size" / "EU Size"
    EAN      : "International Article Number (EAN/UPC)" / "International
               Article Number" / "EAN/UPC" / "EAN" / "UPC" / "Barcode" /
               "Bar Code"
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

import openpyxl

log = logging.getLogger(__name__)

# ══════════════════════════════════════════════════════════════════════════════
# PATHS  (mirrors k_swiss/lambda_function.py's dirs — "fob_order_info" bucket)
# ══════════════════════════════════════════════════════════════════════════════
BASE_DIR = Path(os.environ.get("LAMBDA_TMP_DIR", str(Path(__file__).resolve().parent)))

EAN_DIR     = BASE_DIR / "input" / "fob_order_info"
XML_OUT_DIR = BASE_DIR / "output" / "xml"
LOG_DIR     = BASE_DIR / "output" / "logs"

for _d in (EAN_DIR, XML_OUT_DIR, LOG_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# ══════════════════════════════════════════════════════════════════════════════
# BRAND CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════
BRAND_NAME = "KSwiss"
BRAND_CODE = "KSW"

# ══════════════════════════════════════════════════════════════════════════════
# STIBO / STEP XML CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"

ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z0-9]+)?="[^"]*"')


# ══════════════════════════════════════════════════════════════════════════════
# DATA MODEL
# ══════════════════════════════════════════════════════════════════════════════
@dataclass
class VariantInfo:
    size_raw:    str
    size_code:   str
    ean:         str
    variant_key: str   # KEY_InboundVariant value


@dataclass
class GenericGroup:
    generic_key:  str            # KEY_InboundArticle value
    material_no:  str            # MATERIAL # (article_code, no brand prefix)
    variants:     dict[str, VariantInfo] = field(default_factory=dict)
    # key = size_code (deduplication); last EAN wins on collision


# ══════════════════════════════════════════════════════════════════════════════
# SIZE CODE FORMATTER  (same 3-char MAA convention as implus/dr_marten)
# ══════════════════════════════════════════════════════════════════════════════
_SIZE_OVERRIDES: dict[str, str] = {
    "ONE SIZE": "ONS",
    "ONE-SIZE": "ONS",
    "OS":       "ONS",
    "ONESIZE":  "ONS",
    "FREE":     "ONS",
    "FREE SIZE": "ONS",
    "NS":       "NSZ",
    "NO SIZE":  "NSZ",
}


def _size_code(size_val: str) -> str:
    """
    Convert a raw Europe-size string to a 3-char MAA size code for use in
    KEY_InboundVariant.

    Priority:
      1. Override table (ONE SIZE → ONS, etc.)
      2. Pure integer → zero-padded to 3 digits (e.g. "9" → "009")
      3. Half-size decimal (e.g. "9.5") → "09H"
      4. Other decimal → leading integer zero-padded (e.g. "12.0" → "012")
      5. 1-char alpha → "00X"  (e.g. "M" → "00M")
      6. 2-char alpha → "0XY"  (e.g. "XL" → "0XL")
      7. 3+ char alpha/alnum → first 3 chars uppercased (e.g. "XXL" → "XXL")
    """
    if not size_val:
        return "MSC"
    s = str(size_val).strip().upper()
    if not s:
        return "MSC"

    if s in _SIZE_OVERRIDES:
        return _SIZE_OVERRIDES[s]

    if re.match(r"^\d+$", s):
        try:
            return str(int(s)).zfill(3)[:3]
        except ValueError:
            pass

    m = re.match(r"^(\d+)\.5$", s)
    if m:
        return str(int(m.group(1))).zfill(2)[:2] + "H"

    m2 = re.match(r"^(\d+)\.\d+$", s)
    if m2:
        return str(int(m2.group(1))).zfill(3)[:3]

    alnum = re.sub(r"[^A-Z0-9]", "", s)
    if not alnum:
        return "MSC"
    if len(alnum) == 1:
        return "00" + alnum
    if len(alnum) == 2:
        return "0" + alnum
    return alnum[:3]


# ══════════════════════════════════════════════════════════════════════════════
# STRING HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _s(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "NaT") else s


def _fmt_ean(val) -> str:
    """Return EAN/UPC as a clean integer string (strips .0 from floats)."""
    if val is None:
        return ""
    try:
        f = float(str(val).strip())
        return str(int(f))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s in ("None", "nan") else s


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL READER
# ══════════════════════════════════════════════════════════════════════════════
_GENERIC_ALIASES  = ["generic number", "generic no", "generic #", "generic code"]
_STYLE_ALIASES    = ["principal style code", "principal style", "style code", "style"]
_COLOR_ALIASES    = ["principal color code", "principal colour code", "color code", "colour code", "color", "colour"]
_WIDTH_ALIASES    = ["width"]
# Deliberately excludes "material number" / "material no" — in the K-Swiss
# EAN file that phrasing names a PER-VARIANT code (e.g. "04413-134-M-10"),
# not the Generic key. "Generic Number" (above) is checked first and covers
# that case; this stays a literal-hash fallback for a differently-shaped file.
_MATERIAL_ALIASES = ["material #", "material#"]
_SIZE_ALIASES     = ["size", "europe size", "eu size"]
_EAN_ALIASES      = [
    "international article number (ean/upc)", "international article number",
    "ean/upc", "ean", "upc", "barcode", "bar code", "ean code", "ean-code",
]


def _find_column(headers: list[str], aliases: list[str]) -> Optional[str]:
    lower_map = {h.strip().lower(): h for h in headers if h}
    for alias in aliases:
        if alias in lower_map:
            return lower_map[alias]
    return None


class EANSourceLoader:
    """Loads the K-Swiss EAN/UPC source Excel — style/color/size/EAN rows."""

    def __init__(self, path: Path):
        self.path = path
        self.headers: list[str] = []
        self.rows: list[dict] = []
        self._load()

    def _load(self) -> None:
        log.info("[KSwiss-EAN] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(str(self.path), data_only=True, read_only=True)

        target_sheet = None
        target_hdr_idx = None
        target_headers: list[str] = []

        for sn in wb.sheetnames:
            ws = wb[sn]
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True)):
                candidate = [str(v).strip() if v is not None else "" for v in row]
                if _find_column(candidate, _GENERIC_ALIASES) or _find_column(candidate, _MATERIAL_ALIASES) or (
                    _find_column(candidate, _STYLE_ALIASES) and _find_column(candidate, _SIZE_ALIASES)
                ):
                    target_sheet, target_hdr_idx, target_headers = sn, row_idx, candidate
                    break
            if target_sheet:
                break

        if not target_sheet:
            log.warning("[KSwiss-EAN] No recognisable header found in any sheet — using first sheet, row 1")
            target_sheet, target_hdr_idx = wb.sheetnames[0], 0
            target_headers = [str(v).strip() if v is not None else "" for v in
                               next(wb[target_sheet].iter_rows(min_row=1, max_row=1, values_only=True), [])]

        log.info("[KSwiss-EAN] Using sheet '%s', header row %d: %s", target_sheet, target_hdr_idx + 1, target_headers)
        self.headers = target_headers

        generic_col  = _find_column(target_headers, _GENERIC_ALIASES)
        material_col = _find_column(target_headers, _MATERIAL_ALIASES)
        style_col    = _find_column(target_headers, _STYLE_ALIASES)
        color_col    = _find_column(target_headers, _COLOR_ALIASES)
        width_col    = _find_column(target_headers, _WIDTH_ALIASES)
        size_col     = _find_column(target_headers, _SIZE_ALIASES)
        ean_col      = _find_column(target_headers, _EAN_ALIASES)

        log.info(
            "[KSwiss-EAN] Columns — Generic='%s' Material='%s' Style='%s' Color='%s' Width='%s' Size='%s' EAN='%s'",
            generic_col, material_col, style_col, color_col, width_col, size_col, ean_col,
        )
        if not generic_col and not material_col and not (style_col and color_col):
            log.error("[KSwiss-EAN] Cannot find Generic Number / Material # column, nor a Style+Color pair — no lookup key available")
        if not size_col:
            log.warning("[KSwiss-EAN] Size column NOT FOUND — sizes will default to 'MSC'")
        if not ean_col:
            log.warning("[KSwiss-EAN] EAN/UPC column NOT FOUND — barcodes will be empty")

        col_idx = {h: i for i, h in enumerate(target_headers)}

        def _cell(row, col_name):
            if not col_name:
                return None
            idx = col_idx.get(col_name)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        rows_out: list[dict] = []
        ws = wb[target_sheet]
        for row in ws.iter_rows(min_row=target_hdr_idx + 2, values_only=True):
            if row is None:
                continue

            generic_raw  = _s(_cell(row, generic_col))
            material_raw = _s(_cell(row, material_col))
            style_raw    = _s(_cell(row, style_col))
            color_raw    = _s(_cell(row, color_col))
            width_raw    = _s(_cell(row, width_col))

            if generic_raw:
                article_code = generic_raw
            elif material_raw:
                article_code = material_raw
            elif style_raw and color_raw:
                article_code = f"{style_raw}-{color_raw}"
                if width_raw:
                    article_code = f"{article_code}-{width_raw}"
            else:
                continue  # no usable lookup key on this row

            size_raw = _s(_cell(row, size_col))
            ean_raw  = _fmt_ean(_cell(row, ean_col))

            rows_out.append({
                "article_code": article_code,
                "size":         size_raw,
                "ean":          ean_raw,
            })

        wb.close()
        self.rows = rows_out
        log.info("[KSwiss-EAN] Parsed %d data rows", len(self.rows))


# ══════════════════════════════════════════════════════════════════════════════
# GROUPING — rows → GenericGroup dict
# ══════════════════════════════════════════════════════════════════════════════
def group_rows(rows: list[dict], brand_code: str) -> dict[str, GenericGroup]:
    """
    Group flat EAN rows into GenericGroup objects.

    Generic key (KEY_InboundArticle) : brand_code[:3].upper() + MATERIAL#
        — MUST match global_line_order_form.py's generic_code formula so
        barcodes attach to the correct Generic article.
    Variant key (KEY_InboundVariant) : generic_key + 3-char size code

    On size_code collision within a generic, the last EAN encountered
    wins (same approach as implus/dr_marten EAN sources).
    """
    generics: dict[str, GenericGroup] = {}
    brand_prefix = brand_code[:3].upper()

    for row in rows:
        material_no = row["article_code"]
        generic_key = f"{brand_prefix}{material_no}"
        sc          = _size_code(row["size"])
        variant_key = f"{generic_key}{sc}"

        if generic_key not in generics:
            generics[generic_key] = GenericGroup(
                generic_key=generic_key,
                material_no=material_no,
            )

        g = generics[generic_key]
        g.variants[sc] = VariantInfo(
            size_raw=row["size"],
            size_code=sc,
            ean=row["ean"],
            variant_key=variant_key,
        )

    log.info(
        "[KSwiss-EAN] %d rows → %d generics, %d variants total",
        len(rows), len(generics), sum(len(g.variants) for g in generics.values()),
    )
    return generics


# ══════════════════════════════════════════════════════════════════════════════
# XML HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _add_barcode_datacontainer(parent: ET.Element, ean: str) -> None:
    """
    Append DC_Barcode DataContainer to a Variant product element.

        <DataContainers>
          <MultiDataContainer Type="DC_Barcode">
            <DataContainer Analyzer="true">
              <Values>
                <Value AttributeID="AT_Barcode">190665653885</Value>
                <Value AttributeID="AT_BarcodeType" ID="P"/>
                <Value AttributeID="AT_MainEANIndicator" ID="Y"/>
              </Values>
            </DataContainer>
          </MultiDataContainer>
        </DataContainers>
    """
    if not ean:
        return

    dc_root = ET.SubElement(parent, f"{{{STIBO_NS}}}DataContainers")
    mdc     = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc      = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv     = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = ean

    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


def _build_generic_xml(group: GenericGroup) -> str:
    """
    Build XML string for one Generic article with all its Variants.

        <Product UserTypeID="PRD_GenericArticle">
          <KeyValue KeyID="KEY_InboundArticle">KSW04437-031-M</KeyValue>
          <Values/>
          <Product UserTypeID="PRD_VariantArticle">
            <KeyValue KeyID="KEY_InboundVariant">KSW04437-031-M009</KeyValue>
            <Values/>
            <DataContainers>...</DataContainers>
          </Product>
          ...
        </Product>
    """
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = group.generic_key

    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

    for variant in group.variants.values():
        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant.variant_key

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

        _add_barcode_datacontainer(v_el, variant.ean)

    return ET.tostring(g_el, encoding="unicode")


# ══════════════════════════════════════════════════════════════════════════════
# XML WRITER
# ══════════════════════════════════════════════════════════════════════════════
def build_xml(generics: dict[str, GenericGroup], out_path: Path) -> tuple[int, int]:
    export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = 0
    variant_count = 0

    # ══════════════════════════════════════════════════════════════════════
    # TEST LIMITER — Set TEST_MODE = False for production
    # Caps output to PRODUCT_LIMIT generics, each truncated to at most
    # VARIANT_PER_PRODUCT_LIMIT of its variants (not all 12-15 real sizes).
    # ══════════════════════════════════════════════════════════════════════
    TEST_MODE = False
    PRODUCT_LIMIT = 1
    VARIANT_PER_PRODUCT_LIMIT = 3

    log.info("[KSwiss-EAN] Writing STEP XML → %s", out_path.name)
    with open(out_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n'
        )
        f.write("  <Products>\n")

        for group in generics.values():
            if not group.variants:
                continue
            if TEST_MODE:
                if generic_count >= PRODUCT_LIMIT:
                    break
                truncated_variants = dict(list(group.variants.items())[:VARIANT_PER_PRODUCT_LIMIT])
                group = GenericGroup(
                    generic_key=group.generic_key,
                    material_no=group.material_no,
                    variants=truncated_variants,
                )
            xml_str = _build_generic_xml(group)
            xml_str = _XMLNS_RE.sub("", xml_str)
            f.write(f"    {xml_str}\n")
            generic_count += 1
            variant_count += len(group.variants)

        f.write("  </Products>\n")
        f.write("</STEP-ProductInformation>\n")

    if TEST_MODE:
        log.info(
            "[KSwiss-EAN] TEST MODE: limited to %d product(s), %d variant(s) each",
            PRODUCT_LIMIT, VARIANT_PER_PRODUCT_LIMIT,
        )

    log.info(
        "[KSwiss-EAN] Done — %d generics, %d variants written → %s",
        generic_count, variant_count, out_path.name,
    )
    return generic_count, variant_count


# ══════════════════════════════════════════════════════════════════════════════
# ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
def run(args, auditor=None):
    """
    Main entry point — called by k_swiss/lambda_function.py's ETL_DISPATCHER
    for the "fob_order_info" file type (K-Swiss EAN/UPC source upload,
    e.g. "FW26 UPC - PT.MAP 20260508.xlsx").

    args attributes used:
        brand_code   (str)  default "KSW"
        input_file   (str)  optional: explicit path to the EAN Excel
    """
    brand_code = getattr(args, "brand_code", BRAND_CODE) or BRAND_CODE
    log.info("[KSwiss-EAN] Starting — brand_code=%s", brand_code)

    input_file_arg = getattr(args, "input_file", None)
    if input_file_arg and Path(input_file_arg).exists():
        ean_file = Path(input_file_arg)
    else:
        candidates = sorted(
            (p for p in EAN_DIR.iterdir() if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")),
            key=lambda p: p.stat().st_mtime,
        ) if EAN_DIR.exists() else []
        if not candidates:
            log.warning("[KSwiss-EAN] No EAN .xlsx/.xlsm found in %s — nothing to process.", EAN_DIR)
            return [], None
        ean_file = candidates[-1]

    log.info("[KSwiss-EAN] Processing file: %s", ean_file.name)

    loader = EANSourceLoader(ean_file)
    if not loader.rows:
        log.warning("[KSwiss-EAN] No valid rows parsed from %s", ean_file.name)
        return [], None

    generics = group_rows(loader.rows, brand_code)
    if not generics:
        log.warning("[KSwiss-EAN] No valid generics after grouping.")
        return [], None

    xml_path = XML_OUT_DIR / f"{ean_file.stem}_EAN.xml"
    generic_count, variant_count = build_xml(generics, xml_path)

    print(
        f"=== K-SWISS INLINE EAN SUMMARY ===\n"
        f"  Input rows : {len(loader.rows)}\n"
        f"  Generics   : {generic_count}\n"
        f"  Variants   : {variant_count}\n"
        f"  XML file   : {xml_path.name}",
        flush=True,
    )

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return [], xml_path


# ══════════════════════════════════════════════════════════════════════════════
# CLI (for local testing)
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )

    ap = argparse.ArgumentParser(description="K-Swiss Inline EAN → STEP XML")
    ap.add_argument("--brand-code", default=BRAND_CODE, dest="brand_code")
    ap.add_argument("--input-file", default=None,        dest="input_file",
                     help="Path to K-Swiss EAN/UPC source Excel (optional)")
    run(ap.parse_args())
