"""
╔══════════════════════════════════════════════════════════════════╗
║   STIBO INBOUND XML GENERATOR — Birkenstock EAN  v2.0            ║
║   Birkenstock EAN report → Stibo STEP XML (Generic + Variants)   ║
║   (Structure mirrors crocs/order_form_fw.py)                     ║
╚══════════════════════════════════════════════════════════════════╝

Input file: Birkenstock EAN report (Excel)
    Sheet   : DATA
    Style   : Material            (e.g. "1460-BKCH")
    Size    : Grid Value          (column K)
    Barcode : EAN / Barcode

KEY_InboundArticle formula:
    BrandCode + Material   (full value, e.g. BCK1460-BKCH)

KEY_InboundVariant formula:
    KEY_InboundArticle + Size LOV ID (from MDD "Size Code LOV" lookup,
    Col G description → Col F LOV ID; fallback formula if not found)

XML shape mirrors Crocs Order Form FW: Generic <Values/> empty,
Variant <Values/> empty, barcode goes into DC_Barcode DataContainer.
"""

import re
import sys
import logging
import argparse
import os
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

# ======================================================================
# PATHS
# ======================================================================
BASE_DIR = Path(os.environ.get("LAMBDA_TMP_DIR", Path(__file__).resolve().parent))

EAN_DIR     = BASE_DIR / "input" / "oor"
MDD_DIR     = BASE_DIR / "input" / "mdd"
ATTR_DIR    = BASE_DIR / "input" / "attributes"
XML_OUT_DIR = BASE_DIR / "output" / "xml"
LOG_DIR     = BASE_DIR / "output" / "logs"

for _d in (EAN_DIR, MDD_DIR, ATTR_DIR, XML_OUT_DIR, LOG_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# ======================================================================
# BRAND CONSTANTS
# ======================================================================
BRAND_NAME = "Birkenstock"
BRAND_CODE = "BCK"

# Test mode: cap to N products, M variants per product (applied after grouping)
# TEST_MODE = True
# TEST_MAX_PRODUCTS = 5
# TEST_MAX_VARIANTS_PER_PRODUCT = 2
# ======================================================================
# XM L NAMESPACE
# ======================================================================
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"
ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z]+)?="[^"]*"')


# ======================================================================
# SIZE FALLBACK FORMULA (used only if MDD Size Code LOV lookup misses)
# ======================================================================

def _maa_size_code_fallback(size_val: str) -> str:
    """Fallback 3-char MAA size code when the value isn't in the MDD
    'Size Code LOV'. Same rule shape as the Crocs fallback formula:
      10    -> 010   (integer numeric, zero-padded to 3)
      6.5   -> 06H   (half-size)
      M     -> 00M   (1-char alpha)
      XL    -> 0XL   (2-char alpha)
      ABC   -> ABC   (3+ char, first 3 as-is)
    """
    s = size_val.strip().upper()
    if not s:
        return "MSC"

    if re.match(r'^\d+$', s):
        try:
            return str(int(s)).zfill(3)[:3]
        except ValueError:
            pass

    m = re.match(r'^(\d+)\.5$', s)
    if m:
        return str(int(m.group(1))).zfill(2)[:2] + "H"

    m2 = re.match(r'^(\d+)\.\d+$', s)
    if m2:
        return str(int(m2.group(1))).zfill(3)[:3]

    alnum = re.sub(r"[^A-Z0-9]", "", s)
    if not alnum:
        return "MSC"
    if len(alnum) == 1:
        return "00" + alnum
    if len(alnum) == 2:
        return "0" + alnum
    return alnum[:3]


def _maa_size_code(size_val: str, size_lov: dict | None = None) -> str:
    """3-char MAA size code for use in KEY_InboundVariant.

    Priority:
      1. MDD 'Size Code LOV' lookup (Col G description -> Col F LOV ID).
      2. Fallback formula.
    """
    s = size_val.strip().upper()

    if size_lov:
        if s in size_lov:
            return size_lov[s]
        stripped = s.lstrip("0") or "0"
        if stripped in size_lov:
            return size_lov[stripped]

    return _maa_size_code_fallback(s)


# ======================================================================
# LOADERS
# ======================================================================

class MDDLoader:
    """Loads Core Attributes + LOVs from the MDD Excel."""

    def __init__(self, path: Path):
        self.path = path
        self.core: list[dict] = []
        self.lovs: dict[str, dict] = {}
        self._load()

    def _load(self):
        log.info("[MDD] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        if "Core Attributes" in wb.sheetnames:
            ws   = wb["Core Attributes"]
            rows = list(ws.iter_rows(values_only=True))
            hdr  = rows[1] if len(rows) > 1 else rows[0]
            col  = {str(h).strip(): i for i, h in enumerate(hdr) if h}
            for row in rows[2:]:
                attr_id = row[7] if len(row) > 7 else None
                if not attr_id:
                    continue
                attr_id = str(attr_id).strip()
                def _v(key, _row=row, _col=col):
                    idx = _col.get(key)
                    if idx is None or idx >= len(_row): return None
                    v = _row[idx]; return str(v).strip() if v else None
                self.core.append({
                    "id": attr_id,
                    "name": _v("PIM Attribute Name"),
                    "cardinality": _v("Cardinality"),
                    "lov_name": _v("Name of LOV")
                })
        self._load_simple_lovs(wb)
        self._load_named_lov_sheets(wb)
        self._load_size_lov(wb)   # Col G (description) -> Col F (LOV ID)
        wb.close()
        log.info("[MDD] %d attributes | %d LOVs", len(self.core), len(self.lovs))

    def _load_size_lov(self, wb):
        """Load Size Code LOV: Col G = description, Col F = Stibo LOV ID.
        Stored as lovs['Size Code'] = { description: lov_id }."""
        if "Size Code LOV" not in wb.sheetnames:
            return
        rows = list(wb["Size Code LOV"].iter_rows(values_only=True))
        size_map: dict = {}
        for row in rows[1:]:
            if not row or len(row) < 7:
                continue
            lov_id = row[5]   # Col F
            desc   = row[6]   # Col G
            if lov_id is None or desc is None:
                continue
            lid = str(lov_id).strip()
            key = str(desc).strip().upper()
            if key.endswith('.'): key = key[:-1]
            if key and lid:
                size_map.setdefault(key, lid)
        if size_map:
            self.lovs["Size Code"] = size_map
            log.info("[MDD] Size Code LOV (colF/G): %d entries", len(size_map))

    def _load_simple_lovs(self, wb):
        if "Simple LOVs" not in wb.sheetnames: return
        for row in list(wb["Simple LOVs"].iter_rows(values_only=True))[1:]:
            lov_name, _, val_name, val_id = (row + (None,)*4)[:4]
            if lov_name and val_name:
                self.lovs.setdefault(str(lov_name).strip(), {})[str(val_name).strip()] = (
                    str(val_id).strip() if val_id else str(val_name).strip())

    def _load_named_lov_sheets(self, wb):
        for sn in wb.sheetnames:
            if "LOV" not in sn.upper(): continue
            rows = list(wb[sn].iter_rows(values_only=True))
            if len(rows) < 2: continue
            display = sn.replace(" LOV","").replace("LOV ","").strip()
            for row in rows[1:]:
                if not row or len(row) < 2: continue
                code, name = row[0], row[1]
                if name:
                    self.lovs.setdefault(display, {})[str(name).strip()] = (
                        str(code).strip() if code else str(name).strip())


class BCKEANLoader:
    """Loads Birkenstock EAN report source data from the 'DATA' sheet."""

    SHEET_NAME = "DATA"

    def __init__(self, path: Path):
        self.path = path
        self.df = pd.DataFrame()
        self.headers: list[str] = []
        self._load()

    def _load(self):
        log.info("[Birkenstock-EAN] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)

        target = None
        for sn in wb.sheetnames:
            if sn.strip().upper() == self.SHEET_NAME:
                target = sn
                break
        if not target:
            log.warning("[Birkenstock-EAN] Sheet 'DATA' not found — using first sheet")
            target = wb.sheetnames[0]

        log.info("[Birkenstock-EAN] Using sheet: '%s'", target)
        ws = wb[target]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            log.warning("[Birkenstock-EAN] Sheet '%s' is empty", target)
            return

        # Header row = first row with a "Material" column
        hdr_idx = 0
        for i in range(min(15, len(rows))):
            if rows[i]:
                rv = {str(v).strip().lower() for v in rows[i] if v is not None}
                if "material" in rv:
                    hdr_idx = i
                    break

        log.info("[Birkenstock-EAN] Header at row %d (0-indexed)", hdr_idx)
        header = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[hdr_idx])]

        seen_cols: dict = {}
        for ci, col in enumerate(header):
            if col in seen_cols:
                seen_cols[col] += 1
                header[ci] = f"{col}_{seen_cols[col]}"
            else:
                seen_cols[col] = 0

        df = pd.DataFrame(rows[hdr_idx + 1:], columns=header)

        style_col = self._find_column(header, ["Material"])
        if style_col:
            df = df[df[style_col].notna() & (~df[style_col].astype(str).str.strip().isin(["", "None", "nan"]))]

        self.df = df
        self.headers = header

        log.info("[Birkenstock-EAN] %d rows loaded", len(self.df))
        log.info("[Birkenstock-EAN] Headers: %s", self.headers[:15])

    def _find_column(self, headers: list, candidates: list) -> str | None:
        for c in candidates:
            if c in headers:
                return c
            for h in headers:
                if h.lower() == c.lower():
                    return h
        return None

    def find_column(self, *candidates) -> str | None:
        return self._find_column(self.headers, list(candidates))

    def column_by_letter(self, letter: str) -> str | None:
        """Return the header name sitting at a specific Excel column letter
        (e.g. 'K' -> index 10). Used when duplicate header names (e.g. two
        'Grid Value' columns) make name-based lookup ambiguous."""
        idx = 0
        for ch in letter.strip().upper():
            idx = idx * 26 + (ord(ch) - ord("A") + 1)
        idx -= 1  # convert to 0-based
        if 0 <= idx < len(self.headers):
            return self.headers[idx]
        return None


# ======================================================================
# HELPERS
# ======================================================================

def _s(v) -> str:
    if v is None: return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "0", "NaT") else s


def _fmt_ean(ean_val) -> str:
    if ean_val is None: return ""
    try:
        return str(int(float(str(ean_val).strip())))
    except (ValueError, TypeError):
        s = str(ean_val).strip()
        return "" if s in ("None", "nan") else s


def _clean_style(style_val) -> str:
    if style_val is None: return ""
    s = str(style_val).strip()
    return "" if s in ("None", "nan", "") else s


def _parse_brand_code_from_filename(path: Path, default: str = BRAND_CODE) -> str:
    """
    Extract brand code from a filename segment like:
      'BIRKENSTOCK (BCX)'  -> BCX
      'BIRKENSTOCK BCX'    -> BCX
    Falls back to `default` if no brand code token is found.
    """
    stem = path.stem.upper()

    # 1. Parenthesised code, e.g. BIRKENSTOCK (BCX)
    match = re.search(r"\(([A-Z0-9]{2,6})\)", stem)
    if match:
        return match.group(1)

    # 2. Space-separated code right after BIRKENSTOCK, e.g. BIRKENSTOCK BCX
    match = re.search(r"BIRKENSTOCK\s+([A-Z0-9]{2,6})(?:[\s\-]|$)", stem)
    if match:
        return match.group(1)

    return default


# ======================================================================
# MAPPER — groups flat rows into Generic + Variants dict
# ======================================================================

def group_ean_rows(raw_rows: list[dict], brand_code: str,
                    loader: BCKEANLoader, mdd=None) -> dict[str, dict]:
    """
    KEY_InboundArticle = BrandCode + Material (full value, e.g. BCK1460-BKCH)
    KEY_InboundVariant = KEY_InboundArticle + Size LOV ID

    Each unique Material = One Generic.
    Each row = One Variant (Grid Value = size, EAN = barcode).
    """
    result: dict[str, dict] = {}

    style_col = loader.find_column("Material")
    # Grid Value / size lives specifically at column K — some sheets have a
    # second "Grid Value"-named column elsewhere, so pick by position, not name.
    size_col  = loader.column_by_letter("K")
    ean_col   = loader.find_column("EAN-Code", "EAN Code", "EAN", "Barcode", "Bar code", "EAN13", "EAN/UPC", "UPC")

    if not style_col:
        log.error("[Birkenstock-EAN] Cannot find Material column")
        return result

    log.info("[Birkenstock-EAN] All available headers: %s", loader.headers)
    log.info("[Birkenstock-EAN] Using columns: Style='%s', Size(col K)='%s', EAN='%s'",
             style_col, size_col, ean_col)

    if not size_col:
        log.warning("[Birkenstock-EAN] Grid Value column NOT FOUND - sizes will default to 'MSC'!")
    if not ean_col:
        log.warning("[Birkenstock-EAN] EAN/Barcode column NOT FOUND - barcodes will be empty!")

    size_lov = mdd.lovs.get("Size Code") if mdd else None
    if size_lov:
        log.info("[Birkenstock-EAN] Size Code LOV loaded with %d entries", len(size_lov))
    else:
        log.warning("[Birkenstock-EAN] Size Code LOV not found in MDD - using fallback formulas")

    for row in raw_rows:
        style_raw = _clean_style(row.get(style_col))
        if not style_raw:
            continue

        size_raw = _s(row.get(size_col)) if size_col else ""
        sap_size = _maa_size_code(size_raw, size_lov) if size_raw else "MSC"

        ean_str = _fmt_ean(row.get(ean_col)) if ean_col else ""

        generic_code = f"{brand_code}{style_raw}"

        if generic_code not in result:
            result[generic_code] = {
                "article_no": style_raw,
                "style_code": style_raw,
                "generic_code": generic_code,
                "brand_code": brand_code,
                "division": "F",  # Footwear default for Birkenstock
                "variants": {},
            }

        variant_code = f"{generic_code}{sap_size}"
        variant_key = sap_size

        g = result[generic_code]
        if variant_key not in g["variants"]:
            g["variants"][variant_key] = {
                "size_raw": size_raw,
                "sap_size": sap_size,
                "ean": ean_str,
                "variant_code": variant_code,
            }

    log.info("[Mapper] %d rows -> %d generics  %d variants",
             len(raw_rows), len(result), sum(len(g["variants"]) for g in result.values()))
    return result


# ======================================================================
# XML HELPERS
# ======================================================================

def _add_barcode_datacontainer(parent_el: ET.Element, ean: str) -> None:
    """Add DC_Barcode DataContainer with barcode attributes."""
    if not ean:
        return
    dc_root = ET.SubElement(parent_el, f"{{{STIBO_NS}}}DataContainers")
    mdc = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = ean

    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


# ======================================================================
# PRODUCT XML BUILDER
# ======================================================================

def build_product_xml(generic: dict) -> str:
    """Build XML for one Generic article with its Variants."""
    if not generic.get("article_no"):
        return ""

    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = generic["generic_code"]

    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")

    for var_key in sorted(generic["variants"].keys()):
        variant = generic["variants"][var_key]

        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant["variant_code"]

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")

        if variant.get("ean"):
            _add_barcode_datacontainer(v_el, variant["ean"])

    return ET.tostring(g_el, encoding="unicode")


# ======================================================================
# ORCHESTRATOR
# ======================================================================

def run(args, auditor=None):
    """Main entry point for Birkenstock EAN processing."""
    brand      = getattr(args, "brand", BRAND_NAME)
    brand_code = getattr(args, "brand_code", BRAND_CODE)

    log.info("[Birkenstock-EAN] Starting: brand=%s code=%s", brand, brand_code)

    mdd = None
    mdd_file = next(MDD_DIR.glob("*.xlsx"), None)
    if mdd_file:
        mdd = MDDLoader(mdd_file)
    else:
        log.warning("[Birkenstock-EAN] No MDD file found in %s", MDD_DIR)

    ean_file = next(EAN_DIR.glob("*.xls*"), None)
    if not ean_file:
        ean_file = next(EAN_DIR.glob("*.csv"), None)
    if not ean_file:
        raise FileNotFoundError(f"No EAN file in {EAN_DIR}")

    file_brand_code = _parse_brand_code_from_filename(ean_file, default=brand_code)
    if file_brand_code != brand_code:
        log.info("[Birkenstock-EAN] brand_code from filename '%s' overrides arg default '%s'",
                  file_brand_code, brand_code)
    brand_code = file_brand_code

    loader = BCKEANLoader(ean_file)
    if loader.df.empty:
        log.warning("[Birkenstock-EAN] No data rows found")
        return

    raw_rows = loader.df.to_dict("records")
    generics = group_ean_rows(raw_rows, brand_code, loader, mdd=mdd)

    if not generics:
        log.warning("[Birkenstock-EAN] No valid generics produced")
        return

    # ── Test mode: cap to N products, M variants per product ─────────
    # if TEST_MODE:
    #     capped: dict[str, dict] = {}
    #     for gk in list(generics.keys())[:TEST_MAX_PRODUCTS]:
    #         g = generics[gk]
    #         capped_variants = dict(list(g["variants"].items())[:TEST_MAX_VARIANTS_PER_PRODUCT])
    #         capped[gk] = {**g, "variants": capped_variants}
    #     generics = capped
    #     log.info("[TEST MODE] Limited to %d products x max %d variants each",
    #              len(generics), TEST_MAX_VARIANTS_PER_PRODUCT)

    xml_filename = f"{ean_file.stem}_BIRKENSTOCK_EAN.xml"
    xml_path     = XML_OUT_DIR / xml_filename
    export_time  = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = variant_count = 0

    log.info("[Birkenstock-EAN] Writing XML -> %s", xml_filename)
    with open(xml_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n\n'
        )
        f.write("  <Products>\n")

        for generic in generics.values():
            px = build_product_xml(generic)
            if px:
                f.write(f"    {_XMLNS_RE.sub('', px)}\n")
                generic_count += 1
                variant_count += len(generic["variants"])

        f.write("  </Products>\n</STEP-ProductInformation>\n")

    file_kb = xml_path.stat().st_size // 1024
    log.info("[Birkenstock-EAN] XML written: %s  (%d KB)", xml_path.name, file_kb)
    print(f"=== BIRKENSTOCK EAN SUMMARY ===\n"
          f"  Input rows : {len(raw_rows)}\n"
          f"  Generics   : {generic_count}\n"
          f"  Variants   : {variant_count}\n"
          f"  XML size   : {file_kb} KB", flush=True)

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return xml_path


# ======================================================================
# MAIN
# ======================================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Birkenstock EAN -> Stibo XML")
    parser.add_argument("--brand",      default=BRAND_NAME)
    parser.add_argument("--brand_code", default=BRAND_CODE)
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)]
    )
    run(args)