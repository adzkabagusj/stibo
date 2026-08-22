"""
╔══════════════════════════════════════════════════════════════════╗
║   STIBO INBOUND XML GENERATOR — Reebok EAN Source v1.0           ║
║   "Article Attributes EAN Source" → Stibo STEP XML                ║
║   (Barcode data only — companion to article_master_apparel.py /   ║
║    article_master_footwear.py, which send the Generic-level       ║
║    article data with no barcodes)                                 ║
╚══════════════════════════════════════════════════════════════════╝

Source file: Reebok uploads a dedicated "EAN Source" file, own S3 key /
own trigger — e.g.
    "0888-FF-REEBOK-REEBOK_Article Attributes EAN Source
     (MAA Fashion Footwear) Inline-Multi-SM2027-ID-1.xlsx"
— see reebok/lambda_function.py's "ean" file type. Read from its own
input/ean directory, NOT shared with article_master_apparel.py /
article_master_footwear.py's input/apparel or input/footwear.

Sheet: "NuORDER Order Data" (same layout/header-detection convention as
article_master_apparel.py / article_master_footwear.py's ArticleMasterLoader
— one row per SKU/size).

Column mapping (per user direction, 2026-08-19):
    Principal Style Code -> "Article Number"   (join-key source)
    Size                 -> "Size"              (looked up in the MDD
                                                  "Size Code LOV" sheet,
                                                  Col G description -> Col F
                                                  LOV ID)
    Barcode 1             -> "SKU"               — main/piece barcode
                                                    (AT_MainEANIndicator="Y")
    Barcode 2             -> "UPC"               — secondary barcode
                                                    (AT_MainEANIndicator="N")

Both barcodes are written as separate <DataContainer> entries inside the
SAME <MultiDataContainer Type="DC_Barcode"> — same two-entry shape as
PTP/ean_source.py and implus/ean_update_source_harbinger.py.

KEY_InboundArticle formula (must match article_master_apparel.py /
article_master_footwear.py's own generic key so these barcodes attach to
the correct Generic articles):
    brand_code + Article Number     e.g. "REE12345678"

KEY_InboundVariant formula:
    KEY_InboundArticle + Size LOV ID (from MDD "Size LOV" lookup). No
    fallback formula — if the raw Size value isn't found in the MDD sheet,
    that row's variant is skipped entirely rather than guessed.

XML shape mirrors PTP/implus EAN sources: Generic <Values/> empty, Variant
<Values/> empty, barcodes go into a DC_Barcode DataContainer only.
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

# ======================================================================
# PATHS  (own dedicated upload — input/ean, NOT input/apparel or
# input/footwear; see reebok/lambda_function.py's "ean" file type)
# ======================================================================
BASE_DIR = Path(os.environ.get("LAMBDA_TMP_DIR", Path(__file__).resolve().parent))

INPUT_DIR   = BASE_DIR / "input" / "ean"
MDD_DIR     = BASE_DIR / "input" / "mdd"
XML_OUT_DIR = BASE_DIR / "output" / "xml"
LOG_DIR     = BASE_DIR / "output" / "logs"

for _d in (INPUT_DIR, MDD_DIR, XML_OUT_DIR, LOG_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# ======================================================================
# BRAND CONSTANTS
# ======================================================================
BRAND_NAME = "Reebok"
BRAND_CODE = "REE"   # 3-letter SAP/MDD brand code for Reebok (same as article master files)

# ======================================================================
# XML NAMESPACE
# ======================================================================
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"
ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z0-9]+)?="[^"]*"')


# ======================================================================
# STRING HELPERS
# ======================================================================

def _s(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "0", "NaT", "#VALUE!") else s


def _fmt_barcode(val) -> str:
    """
    Return a barcode value as a clean string.

    Only runs the float→int cleanup (stripping a trailing ".0") on actual
    numeric cell values; string cells are returned as-is so significant
    leading zeros (e.g. a text-formatted SKU/UPC) survive.
    """
    if val is None:
        return ""
    if isinstance(val, str):
        s = val.strip()
        return "" if s in ("None", "nan", "NaT") else s
    try:
        f = float(val)
        return str(int(f))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s in ("None", "nan") else s


# ======================================================================
# MDD LOADER  (Size LOV only — this file emits barcode data alone, no
# other attribute needs MDD/RNA lookups; see module docstring)
# ======================================================================

class MDDLoader:
    """Loads the 'Size Code LOV' sheet from the shared Master Data Dictionary Excel."""

    def __init__(self, path: Path):
        self.path = path
        self.lovs: dict[str, dict] = {}
        self._load()

    def _load(self) -> None:
        log.info("[MDD] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        self._load_size_lov(wb)
        wb.close()
        log.info("[MDD] %d LOV types loaded", len(self.lovs))

    def _load_size_lov(self, wb) -> None:
        """Load 'Size Code LOV': Col F (index 5) = Stibo LOV ID, Col G
        (index 6) = description — same sheet + column layout as
        crocs/shipment_confirmation_main.py's MDDLoader._load_size_lov."""
        if "Size Code LOV" not in wb.sheetnames:
            log.warning("[MDD] 'Size Code LOV' sheet not found — no Size LOV lookups will resolve; all rows will be skipped")
            return
        rows = list(wb["Size Code LOV"].iter_rows(values_only=True))
        lov: dict[str, str] = {}
        for row in rows[1:]:
            if not row or len(row) < 7:
                continue
            lov_id = row[5]   # Col F
            desc   = row[6]   # Col G
            if lov_id is None or desc is None:
                continue
            lid = str(lov_id).strip()
            key = str(desc).strip().upper()
            if key.endswith("."):
                key = key[:-1]   # strip trailing dot
            if key and lid:
                lov.setdefault(key, lid)   # first occurrence wins
        self.lovs["SizeLOV"] = lov
        log.info("[MDD] Size Code LOV (colF/G): %d entries", len(lov))


def _resolve_size_lov_id(size_raw: str, size_lov: dict | None) -> str:
    """Resolve a raw Size value to its MDD LOV ID. Returns "" if there's no
    match — no fallback formula; callers must skip the row rather than
    guess a code (see module docstring)."""
    if not size_lov:
        return ""
    s = size_raw.strip().upper()
    if s in size_lov:
        return size_lov[s]
    stripped = s.lstrip("0") or "0"
    if stripped in size_lov:
        return size_lov[stripped]
    return ""


# ======================================================================
# INPUT FILE LOADER  ("NuORDER Order Data" sheet — same header-detection
# convention as article_master_apparel.py / article_master_footwear.py)
# ======================================================================

class EANLoader:
    PREFERRED_SHEETS = ["NuORDER Order Data"]
    HEADER_SIGNAL = "Article Number"

    def __init__(self, path: Path):
        self.path    = path
        self.df      = pd.DataFrame()
        self.headers: list[str] = []
        self._load()

    def _load(self) -> None:
        log.info("[Reebok-EAN] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        log.info("[Reebok-EAN] Available sheets: %s", wb.sheetnames)

        target = next(
            (s for s in wb.sheetnames if s.strip() in self.PREFERRED_SHEETS),
            wb.sheetnames[0],
        )
        log.info("[Reebok-EAN] Using sheet: '%s'", target)

        ws   = wb[target]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            log.warning("[Reebok-EAN] Sheet '%s' is empty", target)
            return

        hdr_idx = 0
        for i in range(min(20, len(rows))):
            if rows[i] and any(
                str(v).strip() == self.HEADER_SIGNAL for v in rows[i] if v is not None
            ):
                hdr_idx = i
                log.info("[Reebok-EAN] Header detected at row %d (found '%s')", hdr_idx, self.HEADER_SIGNAL)
                break
        else:
            log.warning(
                "[Reebok-EAN] '%s' not found in first 20 rows — defaulting to row 0",
                self.HEADER_SIGNAL,
            )

        raw_header = rows[hdr_idx]
        header = [
            str(h).strip() if h is not None else f"col_{i}"
            for i, h in enumerate(raw_header)
        ]

        seen: dict[str, int] = {}
        for ci, col in enumerate(header):
            if col in seen:
                seen[col] += 1
                header[ci] = f"{col}_{seen[col]}"
            else:
                seen[col] = 0

        self.headers = header
        df = pd.DataFrame(rows[hdr_idx + 1:], columns=header)
        df = df.dropna(how="all")

        self.df = df.reset_index(drop=True)
        log.info("[Reebok-EAN] Loaded %d data rows | %d columns", len(self.df), len(header))

    @staticmethod
    def _find_col(headers: list[str], candidates: list[str]) -> str | None:
        for c in candidates:
            if c in headers:
                return c
            for h in headers:
                if h.lower() == c.lower():
                    return h
        return None

    def find_col(self, *candidates: str) -> str | None:
        return self._find_col(self.headers, list(candidates))


# ======================================================================
# DATA MODEL
# ======================================================================

@dataclass
class VariantInfo:
    size_raw:    str
    size_key:    str    # MDD Size LOV ID
    sku_barcode: str
    upc_barcode: str
    variant_key: str    # KEY_InboundVariant value


@dataclass
class GenericGroup:
    generic_key:    str            # KEY_InboundArticle value
    article_number: str
    variants:       dict[str, VariantInfo] = field(default_factory=dict)
    # key = size_key (deduplication); last row wins on collision


# ======================================================================
# GROUPING — rows → GenericGroup dict
# ======================================================================

def group_rows(raw_rows: list[dict], loader: EANLoader, brand_code: str, size_lov: dict | None) -> dict[str, GenericGroup]:
    """
    Group flat per-SKU rows into GenericGroup objects.

    Generic key (KEY_InboundArticle) : brand_code + Article Number
        — matches article_master_apparel.py / article_master_footwear.py's
        own generic key exactly, so these barcodes attach to the correct
        Generic articles those files create.
    Variant key (KEY_InboundVariant) : generic_key + Size LOV ID (MDD
        lookup). Rows whose Size value has no MDD match are skipped — no
        fallback formula, no default guess.
    """
    col_article = loader.find_col("Article Number")
    col_size    = loader.find_col("Size")
    col_sku     = loader.find_col("SKU")
    col_upc     = loader.find_col("UPC")

    log.info(
        "[Reebok-EAN] Columns — Article Number='%s' Size='%s' SKU='%s' UPC='%s'",
        col_article, col_size, col_sku, col_upc,
    )

    if col_article is None:
        log.error("[Reebok-EAN] Cannot find 'Article Number' column. Available columns: %s", loader.headers)
        return {}
    if col_sku is None:
        log.warning("[Reebok-EAN] 'SKU' column NOT FOUND — Barcode 1 will be empty")
    if col_upc is None:
        log.warning("[Reebok-EAN] 'UPC' column NOT FOUND — Barcode 2 will be empty")

    generics: dict[str, GenericGroup] = {}
    skipped_no_size_lov = 0

    for row in raw_rows:
        article_number = _s(row.get(col_article)) if col_article else ""
        if not article_number:
            continue

        generic_key = f"{brand_code}{article_number}"

        size_raw = _s(row.get(col_size)) if col_size else ""
        size_key = _resolve_size_lov_id(size_raw, size_lov) if size_raw else ""
        if size_key and len(size_key) < 3:
            size_key = size_key.zfill(3)
        if not size_key:
            skipped_no_size_lov += 1
            log.info(
                "[Reebok-EAN] Size '%s' not found in MDD Size LOV — skipping variant for %s",
                size_raw, generic_key,
            )
            continue
        variant_key = f"{generic_key}{size_key}"

        sku_barcode = _fmt_barcode(row.get(col_sku)) if col_sku else ""
        upc_barcode = _fmt_barcode(row.get(col_upc)) if col_upc else ""

        if generic_key not in generics:
            generics[generic_key] = GenericGroup(
                generic_key=generic_key,
                article_number=article_number,
            )

        generics[generic_key].variants[size_key] = VariantInfo(
            size_raw=size_raw,
            size_key=size_key,
            sku_barcode=sku_barcode,
            upc_barcode=upc_barcode,
            variant_key=variant_key,
        )

    if skipped_no_size_lov:
        log.warning(
            "[Reebok-EAN] %d row(s) skipped — Size value not found in MDD Size LOV",
            skipped_no_size_lov,
        )

    log.info(
        "[Reebok-EAN] %d rows → %d generics, %d variants total",
        len(raw_rows), len(generics), sum(len(g.variants) for g in generics.values()),
    )
    return generics


# ======================================================================
# XML HELPERS
# ======================================================================

def _add_barcode_datacontainer(parent: ET.Element, sku_barcode: str, upc_barcode: str) -> None:
    """
    Append DC_Barcode DataContainer to a Variant product element — TWO
    <DataContainer> entries inside the same MultiDataContainer:
      Entry 1 (SKU): AT_Barcode=SKU, AT_BarcodeType="P", AT_MainEANIndicator="Y"
      Entry 2 (UPC): AT_Barcode=UPC, AT_BarcodeType="P", AT_MainEANIndicator="N"

        <DataContainers>
          <MultiDataContainer Type="DC_Barcode">
            <DataContainer Analyzer="true">
              <Values>
                <Value AttributeID="AT_Barcode">123456789012</Value>
                <Value AttributeID="AT_BarcodeType" ID="P"/>
                <Value AttributeID="AT_MainEANIndicator" ID="Y"/>
              </Values>
            </DataContainer>
            <DataContainer Analyzer="true">
              <Values>
                <Value AttributeID="AT_Barcode">987654321098</Value>
                <Value AttributeID="AT_BarcodeType" ID="P"/>
                <Value AttributeID="AT_MainEANIndicator" ID="N"/>
              </Values>
            </DataContainer>
          </MultiDataContainer>
        </DataContainers>
    """
    if not sku_barcode and not upc_barcode:
        return

    dc_root = ET.SubElement(parent, f"{{{STIBO_NS}}}DataContainers")
    mdc     = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")

    for barcode, is_main in ((sku_barcode, "Y"), (upc_barcode, "N")):
        if not barcode:
            continue
        dc  = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
        dc.set("Analyzer", "true")
        dcv = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

        v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
        v1.set("AttributeID", "AT_Barcode")
        v1.text = barcode

        v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
        v2.set("AttributeID", "AT_BarcodeType")
        v2.set("ID", "P")

        v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
        v3.set("AttributeID", "AT_MainEANIndicator")
        v3.set("ID", is_main)


def _build_generic_xml(group: GenericGroup) -> str:
    """
    Build XML string for one Generic article with all its Variants.

        <Product UserTypeID="PRD_GenericArticle">
          <KeyValue KeyID="KEY_InboundArticle">REE12345678</KeyValue>
          <Values/>
          <Product UserTypeID="PRD_VariantArticle">
            <KeyValue KeyID="KEY_InboundVariant">REE12345678042</KeyValue>
            <Values/>
            <DataContainers>...</DataContainers>
          </Product>
          ...
        </Product>
    """
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = group.generic_key

    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

    for variant in group.variants.values():
        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant.variant_key

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

        _add_barcode_datacontainer(v_el, variant.sku_barcode, variant.upc_barcode)

    return ET.tostring(g_el, encoding="unicode")


# ======================================================================
# XML WRITER
# ======================================================================

def build_xml(generics: dict[str, GenericGroup], out_path: Path) -> tuple[int, int]:
    export_time = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = 0
    variant_count = 0

    log.info("[Reebok-EAN] Writing STEP XML → %s", out_path.name)
    with open(out_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n'
        )
        f.write("  <Products>\n")

        for group in generics.values():
            if not group.variants:
                continue
            xml_str = _XMLNS_RE.sub("", _build_generic_xml(group))
            f.write(f"    {xml_str}\n")
            generic_count += 1
            variant_count += len(group.variants)

        f.write("  </Products>\n")
        f.write("</STEP-ProductInformation>\n")

    log.info(
        "[Reebok-EAN] Done — %d generics, %d variants written → %s",
        generic_count, variant_count, out_path.name,
    )
    return generic_count, variant_count


# ======================================================================
# FILENAME METADATA PARSER
# ======================================================================

def _parse_metadata_from_filename(dirs: dict, triggered_file_type: str = None) -> dict:
    """
    Parse comp_code, sbu, brand, brand_code, season, seq from the
    triggered EAN file's filename.

    Expected format (dash-separated), same convention as Anta/Crocs:
        {CompanyCode}-{SBU}-{Brand}-{FileType}-{Multi/Mono}-{Season}-{Country}-{Seq}
    e.g. "0888-FF-REEBOK-REEBOK_Article Attributes EAN Source
          (MAA Fashion Footwear) Inline-Multi-SM2027-ID-1.xlsx"
    """
    candidate_dirs = []
    if triggered_file_type and triggered_file_type in dirs:
        candidate_dirs.append(dirs[triggered_file_type])
    candidate_dirs.append(dirs.get("ean"))

    seen: set = set()
    candidate_dirs = [
        d for d in candidate_dirs
        if d is not None and str(d) not in seen and not seen.add(str(d))
    ]

    for folder in candidate_dirs:
        if not folder or not folder.exists():
            continue
        xls_files = list(folder.glob("*.xlsx")) + list(folder.glob("*.xlsm"))
        for xlsx_file in xls_files:
            stem = xlsx_file.stem
            log.info("Attempting metadata parse from filename: '%s'", stem)

            parts = re.split(r"\s*-\s*", stem)
            if len(parts) < 6:
                log.warning(
                    "  Filename '%s' has only %d '-'-separated parts (need ≥6) — skipping",
                    stem, len(parts),
                )
                continue

            comp_code = parts[0].strip()
            sbu       = parts[1].strip()
            if not comp_code:
                log.warning("  comp_code empty — skipping: '%s'", stem)
                continue

            brand = parts[2].strip().title() if len(parts) > 2 else BRAND_NAME

            # ── Locate season token by regex ─────────────────────
            season     = None
            season_idx = None
            for i, p in enumerate(parts):
                if re.match(r"^[A-Z]{2}\d{2,4}$", p.strip(), re.IGNORECASE):
                    season     = p.strip().upper()
                    season_idx = i
                    break

            if not season or season_idx is None:
                log.warning("  No valid season token in '%s' — skipping", stem)
                continue

            trailing = parts[season_idx + 1:]

            seq = 1
            for p in reversed(trailing):
                if p.strip().isdigit():
                    seq = int(p.strip())
                    break

            country_code = ""
            for p in trailing:
                p = p.strip()
                if re.match(r"^[A-Z]{2,3}$", p, re.IGNORECASE) and not p.isdigit():
                    country_code = p.upper()
                    break

            log.info(
                "Parsed: comp_code=%s sbu=%s brand=%s season=%s seq=%s country=%s",
                comp_code, sbu, brand, season, seq, country_code,
            )
            return {
                "comp_code":    comp_code,
                "sbu":          sbu,
                "brand":        brand,
                "brand_code":   BRAND_CODE,
                "season":       season,
                "seq":          seq,
                "country_code": country_code,
            }

    log.warning("Could not parse metadata from filename — using defaults")
    return {
        "comp_code": "0888", "sbu": "SP", "brand": BRAND_NAME,
        "brand_code": BRAND_CODE, "season": "SS27", "seq": 1,
        "country_code": "",
    }


# ======================================================================
# ORCHESTRATOR
# ======================================================================

def run(args, auditor=None):
    """Main entry point — called by reebok/lambda_function.py for the
    'ean' trigger, off its own dedicated 'EAN Source' upload."""
    brand_code = getattr(args, "brand_code", BRAND_CODE) or BRAND_CODE
    log.info("[Reebok-EAN] Starting — brand_code=%s", brand_code)

    mdd = None
    mdd_file = next(MDD_DIR.glob("*.xlsx"), None) or next(MDD_DIR.glob("*.xlsm"), None)
    if mdd_file:
        mdd = MDDLoader(mdd_file)
    else:
        log.warning("[Reebok-EAN] No MDD file found in %s — Size LOV lookups disabled, all rows will be skipped", MDD_DIR)

    input_file_arg = getattr(args, "input_file", None)
    if input_file_arg and Path(input_file_arg).exists():
        ean_file = Path(input_file_arg)
    else:
        ean_file = next(INPUT_DIR.glob("*.xls*"), None)
        if not ean_file:
            log.warning("[Reebok-EAN] No EAN source .xlsx/.xlsm found in %s — nothing to process.", INPUT_DIR)
            return None

    log.info("[Reebok-EAN] Processing file: %s", ean_file.name)

    loader = EANLoader(ean_file)
    if loader.df.empty:
        log.warning("[Reebok-EAN] No data rows found — nothing to process")
        return None

    size_lov = mdd.lovs.get("SizeLOV") if mdd else None
    raw_rows = loader.df.to_dict("records")
    generics = group_rows(raw_rows, loader, brand_code, size_lov)

    if not generics:
        log.warning("[Reebok-EAN] No valid generics produced")
        return None

    xml_path = XML_OUT_DIR / f"{ean_file.stem}.xml"
    generic_count, variant_count = build_xml(generics, xml_path)

    print(
        f"=== REEBOK EAN SUMMARY ===\n"
        f"  Input file : {ean_file.name}\n"
        f"  Input rows : {len(raw_rows)}\n"
        f"  Generics   : {generic_count}\n"
        f"  Variants   : {variant_count}\n"
        f"  XML file   : {xml_path.name}",
        flush=True,
    )

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return xml_path


# ======================================================================
# CLI (local test)
# ======================================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Reebok EAN Source → Stibo XML")
    parser.add_argument("--brand_code", default=BRAND_CODE)
    parser.add_argument("--input-file", default=None, dest="input_file",
                         help="Path to Reebok EAN source Excel (optional)")
    cli_args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )
    run(cli_args)
