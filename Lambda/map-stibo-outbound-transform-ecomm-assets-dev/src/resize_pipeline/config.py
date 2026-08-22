"""
Configuration for the Ecomm Asset Resizing Pipeline.

S3 layout (single bucket)
───────────────────────────
  Image Requirements.xlsx  ->  {S3_BUCKET}/Image Requirements.xlsx
  input XML                ->  {S3_BUCKET}/step-export/ecomm-assets/{filename}.xml
  output ZIPs              ->  {S3_BUCKET}/transformed/ecomm-assets/

STIBO auth
────────────
  OIDC client-credentials flow.
  A module-level token cache is used so the token is reused across warm
  Lambda invocations and only refreshed when near expiry.

Required environment variables
────────────────────────────────
  STIBO_BASE_URL          STIBO host URL
  STIBO_TOKEN_URL         OIDC token endpoint
  STIBO_CLIENT_ID         OIDC client ID
  STIBO_CLIENT_SECRET     OIDC client secret
  S3_BUCKET               Single S3 bucket name (default: map-stibo-outbound-dev)

Optional environment variables
────────────────────────────────
  STIBO_GRANT_TYPE        OAuth grant type      (default: client_credentials)
  STIBO_TOKEN_LEEWAY      Seconds before expiry to refresh token (default: 30)
  STIBO_CONTEXT           STIBO context param   (default: Context1)
  STIBO_WORKSPACE         STIBO workspace param (default: Approved)
"""

import io
import json
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request

import boto3
from openpyxl import load_workbook


# ── STIBO API ─────────────────────────────────────────────────────────────────
ENVIRONMENT_BASE_URL: str     = os.environ.get("STIBO_BASE_URL", "")
STIBO_TOKEN_URL: str          = os.environ.get("STIBO_TOKEN_URL", "")
STIBO_CLIENT_ID: str          = os.environ.get("STIBO_CLIENT_ID", "")
STIBO_CLIENT_SECRET: str      = os.environ.get("STIBO_CLIENT_SECRET", "")
STIBO_GRANT_TYPE: str         = os.environ.get("STIBO_GRANT_TYPE", "client_credentials")
STIBO_TOKEN_LEEWAY: int       = int(os.environ.get("STIBO_TOKEN_LEEWAY", "30"))
STIBO_CONTEXT: str            = os.environ.get("STIBO_CONTEXT", "Context1")
STIBO_WORKSPACE: str          = os.environ.get("STIBO_WORKSPACE", "Approved")
DOWNLOAD_TIMEOUT_SECONDS: int = 30

# ── S3 (single bucket) ───────────────────────────────────────────────────────
S3_BUCKET: str = os.environ.get("S3_BUCKET", "map-stibo-outbound-dev")

# Fixed S3 paths - not configurable via env vars
S3_REQUIREMENTS_KEY: str = "Image Requirements.xlsx"
S3_INPUT_PREFIX: str     = "step-export/ecomm-assets"
S3_OUTPUT_PREFIX: str    = "transformed/ecomm-assets"

# ── JPEG output ───────────────────────────────────────────────────────────────
MAX_FILE_SIZE_BYTES: int  = 500 * 1024
INITIAL_JPEG_QUALITY: int = 95
MIN_JPEG_QUALITY: int     = 10
JPEG_QUALITY_STEP: int    = 5

# ── Runtime dicts (populated at cold-start) ───────────────────────────────────
RESIZE_SPECS:    dict           = {}
CHANNEL_TO_SPEC: dict[str, str] = {}

# ── OIDC token cache ──────────────────────────────────────────────────────────
# Stores { "access_token": str, "expires_at": float (epoch seconds) }
_token_cache: dict = {}


# ── OIDC token helpers ────────────────────────────────────────────────────────

def get_bearer_token() -> str:
    """
    Return a valid STIBO bearer token, fetching a new one only when the
    cached token is missing or within STIBO_TOKEN_LEEWAY seconds of expiry.

    Returns
    -------
    str
        A valid access token string.

    Raises
    ------
    RuntimeError
        If any required OIDC env var is missing or the token request fails.
    """
    # ── Guard: required env vars ──────────────────────────────────────────────
    missing = [
        v for v in ("STIBO_TOKEN_URL", "STIBO_CLIENT_ID", "STIBO_CLIENT_SECRET")
        if not os.environ.get(v)
    ]
    if missing:
        raise RuntimeError(
            f"Missing OIDC environment variable(s): {', '.join(missing)}. "
            "Set them in the Lambda configuration."
        )

    # ── Return cached token if still valid ────────────────────────────────────
    now = time.time()
    if (
        _token_cache.get("access_token")
        and _token_cache.get("expires_at", 0) - now > STIBO_TOKEN_LEEWAY
    ):
        remaining = int(_token_cache["expires_at"] - now)
        print(f"[TOKEN] Reusing cached token - {remaining}s until expiry.")
        return _token_cache["access_token"]

    # ── Fetch new token ───────────────────────────────────────────────────────
    print(f"[TOKEN] Requesting new token from {STIBO_TOKEN_URL} ...")

    encoded_data = urllib.parse.urlencode({
        "grant_type":    STIBO_GRANT_TYPE,
        "client_id":     STIBO_CLIENT_ID,
        "client_secret": STIBO_CLIENT_SECRET,
    }).encode("utf-8")

    req = urllib.request.Request(
        STIBO_TOKEN_URL,
        data=encoded_data,
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        raise RuntimeError(
            f"Token endpoint returned HTTP {exc.code}: "
            f"{exc.read().decode('utf-8')[:200]}"
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Token request to {STIBO_TOKEN_URL} failed: {exc.reason}"
        ) from exc

    access_token = payload.get("access_token")
    expires_in   = int(payload.get("expires_in", 3600))

    if not access_token:
        raise RuntimeError(
            "Token response did not contain 'access_token'. "
            f"Response keys: {list(payload.keys())}"
        )

    # ── Update cache ──────────────────────────────────────────────────────────
    _token_cache["access_token"] = access_token
    _token_cache["expires_at"]   = now + expires_in

    print(f"[TOKEN] New token obtained - expires in {expires_in}s.")
    return access_token


# ── Excel layout constants ────────────────────────────────────────────────────

# Columns D-H (openpyxl 1-based index)
_SPEC_COLUMNS: dict[int, str] = {
    4: "digital_flagship",
    5: "converse",
    6: "footlocker",
    7: "mapclub",
    8: "zalora",
}

# Rows 2-9 contain metadata (Size, DPI, Background, RGB, Color Hex,
# Final Format, Shape, Margin).  The Margin row itself holds no values;
# the 12 margin-data rows begin immediately after it.
_METADATA_ROW_RANGE: tuple[int, int] = (2, 10)   # rows 2-9 inclusive

# Margin groups in order - each group is 4 consecutive rows (Top/Bottom/Left/Right)
# Group 1 -> NW, Group 2 -> AW, Group 3 -> NM and AM (identical values)
_MARGIN_GROUPS: tuple[tuple[str, ...], ...] = (
    ("NW",),
    ("AW",),
    ("NM", "AM"),
)
_MARGIN_SIDES: tuple[str, ...] = ("top", "bottom", "left", "right")

# Hex color pattern: optional # then exactly 6 hex digits
_HEX_COLOR_RE = re.compile(r'^#?([0-9A-Fa-f]{6})$')


# ── Excel helpers ─────────────────────────────────────────────────────────────

def _parse_canvas_size(raw: str) -> tuple[int, int]:
    """
    Parse canvas size from strings like '800 X 800', '800x800', '800 x 800'.

    Returns
    -------
    tuple[int, int]
        (width, height)

    Raises
    ------
    ValueError
        If the string does not match NNN X NNN pattern.
    """
    match = re.match(r'^\s*(\d+)\s*[xX]\s*(\d+)\s*$', str(raw).strip())
    if not match:
        raise ValueError(
            f"Cannot parse canvas size from '{raw}'. "
            "Expected format: '800 X 800' or '800x800'."
        )
    return int(match.group(1)), int(match.group(2))


def _looks_like_hex_color(raw: str) -> bool:
    """Return True if raw looks like a 6-digit hex color (with or without #)."""
    return bool(_HEX_COLOR_RE.match(raw.strip()))


def _find_row_label(ws, row_idx: int) -> str | None:
    """
    Return the first non-empty string value found in columns A, B, or C
    for the given row. Returns None if all three are empty.

    Some sheets place row labels in column B or C rather than A.
    Non-string cells (e.g. numbers) in those columns are skipped.
    """
    for col in (1, 2, 3):
        val = ws.cell(row=row_idx, column=col).value
        if val is not None and isinstance(val, str) and val.strip():
            return val
    return None


# ── Sheet parser ──────────────────────────────────────────────────────────────

def _parse_requirements_sheet(ws) -> tuple[dict, dict[str, str]]:
    """
    Parse the image requirements worksheet.

    Layout
    ------
    Rows 2-9   : metadata headers.
                   Row with label containing 'size'/'dimension' -> canvas size
                   Row with label containing 'dpi'/'resolution' -> DPI
                   Row with label containing 'hex'              -> background color (#RRGGBB)
                   Row with label containing 'margin'           -> marks end of metadata;
                                                                   no data values on this row
    Next 12 rows: margin data in 3 groups of 4 (Top/Bottom/Left/Right)
                   Group 1 -> NW
                   Group 2 -> AW
                   Group 3 -> NM and AM (identical values)
    Remaining  : channel IDs per spec column, one per row, stop at first empty cell

    Returns
    -------
    tuple[dict, dict[str, str]]
        (RESIZE_SPECS, CHANNEL_TO_SPEC)

    Raises
    ------
    ValueError
        If any required spec field is missing or unparseable.
    """
    codes = ("NW", "AW", "NM", "AM")

    # ── Initialise empty spec dicts ───────────────────────────────────────────
    specs: dict = {
        name: {
            "canvas_width":     None,
            "canvas_height":    None,
            "background_color": None,
            "dpi":              None,
            "margins":          {code: {} for code in codes},
        }
        for name in _SPEC_COLUMNS.values()
    }

    # ── Rows 2-9: metadata ────────────────────────────────────────────────────
    # Track which row has the "Margin" label so we know where margin data starts.
    margin_label_row: int | None = None

    for row_idx in range(*_METADATA_ROW_RANGE):

        label_raw = _find_row_label(ws, row_idx)

        # Debug: log all three label columns so layout issues are visible
        col_a_val = ws.cell(row=row_idx, column=1).value
        col_b_val = ws.cell(row=row_idx, column=2).value
        col_c_val = ws.cell(row=row_idx, column=3).value
        print(
            f"  [DEBUG] Row {row_idx}: "
            f"A={col_a_val!r}  B={col_b_val!r}  C={col_c_val!r}  "
            f"label_used={label_raw!r}"
        )

        label = str(label_raw).strip().lower() if label_raw else ""

        # "Margin" row is a section header with no data values in D-H
        if "margin" in label:
            margin_label_row = row_idx
            print(f"  [DEBUG] Margin header found at row {row_idx}")
            continue

        for col_idx, spec_name in _SPEC_COLUMNS.items():
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is None:
                continue
            raw = str(cell_val).strip()
            if not raw:
                continue

            # ── Canvas size (label-based OR value pattern NNN x NNN) ──────────
            if (
                any(kw in label for kw in ("canvas", "size", "dimension"))
                or re.search(r'\d+\s*[xX]\s*\d+', raw)
            ) and specs[spec_name]["canvas_width"] is None:
                try:
                    w, h = _parse_canvas_size(raw)
                    specs[spec_name]["canvas_width"]  = w
                    specs[spec_name]["canvas_height"] = h
                    print(f"  Canvas '{spec_name}': {w}x{h} (row {row_idx})")
                except ValueError as exc:
                    print(f"  [WARN] {exc} - skipped.")

            # ── Background color - ONLY from the "Color Hex" row ─────────────
            # The "Background" row (row 4) contains a text name like "White" and
            # must NOT be used.  Only the row whose label contains "hex" carries
            # the actual #RRGGBB value.
            elif "hex" in label and specs[spec_name]["background_color"] is None:
                if _looks_like_hex_color(raw):
                    # Use the value as-is; it already contains the leading '#'
                    color = raw if raw.startswith("#") else f"#{raw}"
                    specs[spec_name]["background_color"] = color
                    print(f"  BG color '{spec_name}': {color} (row {row_idx})")
                else:
                    print(
                        f"  [DEBUG] Row {row_idx} '{label_raw}' "
                        f"value {raw!r} is not a valid hex color - skipped."
                    )

            # ── DPI (label-based) ─────────────────────────────────────────────
            elif any(kw in label for kw in ("dpi", "ppi", "resolution")):
                try:
                    specs[spec_name]["dpi"] = int(float(raw))
                    print(f"  DPI '{spec_name}': {int(float(raw))} (row {row_idx})")
                except (ValueError, TypeError):
                    print(
                        f"  [WARN] Cannot parse DPI '{raw}' "
                        f"at row {row_idx}, spec '{spec_name}' - skipped."
                    )

    # ── Determine where margin data rows begin ────────────────────────────────
    if margin_label_row is not None:
        margin_start_row = margin_label_row + 1
        print(
            f"  [DEBUG] Margin data starts at row {margin_start_row} "
            f"(label row was {margin_label_row})"
        )
    else:
        margin_start_row = 10   # safe fallback matching known sheet layout
        print(
            f"  [WARN] 'Margin' label not found in rows 2-9. "
            f"Defaulting margin start to row {margin_start_row}."
        )

    # ── Margin rows: 3 groups of 4 (Top / Bottom / Left / Right) ─────────────
    # Group 1 (offsets 0-3)  -> NW
    # Group 2 (offsets 4-7)  -> AW
    # Group 3 (offsets 8-11) -> NM and AM (same values for both codes)
    for group_idx, codes_in_group in enumerate(_MARGIN_GROUPS):
        group_start = margin_start_row + group_idx * 4
        for side_idx, side in enumerate(_MARGIN_SIDES):
            row_idx = group_start + side_idx
            for col_idx, spec_name in _SPEC_COLUMNS.items():
                cell_val = ws.cell(row=row_idx, column=col_idx).value

                if cell_val is None:
                    print(
                        f"  [WARN] Empty margin at row {row_idx}, col {col_idx} "
                        f"(spec='{spec_name}', codes={codes_in_group}, side='{side}') "
                        "- defaulting to 0."
                    )
                    for code in codes_in_group:
                        specs[spec_name]["margins"][code][side] = 0
                    continue

                try:
                    margin_val = int(float(str(cell_val)))
                except (ValueError, TypeError) as exc:
                    raise ValueError(
                        f"Cannot parse margin '{cell_val}' at row {row_idx}, "
                        f"col {col_idx} (spec='{spec_name}', "
                        f"codes={codes_in_group}, side='{side}'): {exc}"
                    ) from exc

                for code in codes_in_group:
                    specs[spec_name]["margins"][code][side] = margin_val

    # ── Channel IDs begin immediately after the 12 margin rows ───────────────
    channel_start_row = margin_start_row + 12
    print(f"  [DEBUG] Channel IDs start at row {channel_start_row}")
    
    _col_d_idx    = 4   # column D = "digital_flagship"
    _sample_vals  = []
    for _r in range(channel_start_row, channel_start_row + 10):
        _raw = ws.cell(row=_r, column=_col_d_idx).value
        _sample_vals.append(_raw)
    print(
        f"  [DEBUG] First 10 raw channel IDs from col D "
        f"(rows {channel_start_row}-{channel_start_row + 9}):\n"
        f"  {[repr(v) for v in _sample_vals]}"
    )

    channel_to_spec: dict[str, str] = {}

    for col_idx, spec_name in _SPEC_COLUMNS.items():
        row_idx   = channel_start_row
        col_count = 0
        while True:
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is None or str(cell_val).strip() == "":
                break

            raw_channel_id = str(cell_val).strip()

            # Normalise: strip trailing -NNNN suffix then uppercase so Excel
            # keys match XML Value IDs regardless of casing or version suffixes.
            channel_id = re.sub(r'-\d{4}$', '', raw_channel_id).upper()

            if channel_id != raw_channel_id:
                print(
                    f"  [DEBUG] Normalised channel ID: "
                    f"{raw_channel_id!r} -> {channel_id!r}"
                )

            channel_to_spec[channel_id] = spec_name
            col_count += 1
            row_idx   += 1

        print(f"  Channels -> '{spec_name}': {col_count} mapping(s)")

    # ── Debug: verify the first 5 normalised keys actually stored ────────────
    _first_5 = list(channel_to_spec.items())[:5]
    print(
        f"  [DEBUG] First 5 CHANNEL_TO_SPEC keys after normalisation:\n"
        + "\n".join(f"    {repr(k)} -> '{v}'" for k, v in _first_5)
    )

    # ── Validate ──────────────────────────────────────────────────────────────
    for name, spec in specs.items():
        if spec["canvas_width"] is None or spec["canvas_height"] is None:
            raise ValueError(
                f"Spec '{name}' is missing canvas size. "
                "Check rows 2-9 for a 'NNN X NNN' cell."
            )
        if spec["background_color"] is None:
            raise ValueError(
                f"Spec '{name}' is missing background_color. "
                "Check rows 2-9 for a row labelled 'Color Hex' "
                "with a cell value in #RRGGBB format."
            )
        if spec["dpi"] is None:
            raise ValueError(
                f"Spec '{name}' is missing DPI. "
                "Check rows 2-9 for a row labelled 'dpi' or 'resolution'."
            )
        for code in codes:
            missing = set(_MARGIN_SIDES) - set(spec["margins"][code].keys())
            if missing:
                raise ValueError(
                    f"Spec '{name}' / code '{code}' missing margin sides: {missing}."
                )

    if not channel_to_spec:
        raise ValueError(
            "No channel IDs found. "
            f"Check that columns D-H have channel IDs from row {channel_start_row} downward."
        )

    return specs, channel_to_spec


# ── Public loader (also callable in unit tests) ───────────────────────────────

def load_config_from_excel(excel_bytes: bytes) -> tuple[dict, dict[str, str]]:
    """
    Parse raw Excel bytes into (RESIZE_SPECS, CHANNEL_TO_SPEC).

    Sheet selection: first sheet whose name contains 'image', 'requirement',
    'spec', or 'resize'; falls back to the first sheet.

    Parameters
    ----------
    excel_bytes : bytes
        Raw content of Image Requirements.xlsx.

    Returns
    -------
    tuple[dict, dict[str, str]]
        (RESIZE_SPECS, CHANNEL_TO_SPEC)
    """
    wb = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)

    sheet_name = wb.sheetnames[0]
    for name in wb.sheetnames:
        if any(kw in name.lower() for kw in ("image", "requirement", "spec", "resize")):
            sheet_name = name
            break

    print(f"[CONFIG] Parsing sheet: '{sheet_name}'")
    return _parse_requirements_sheet(wb[sheet_name])


# ── S3 initialiser (Lambda cold-start) ───────────────────────────────────────

def _init_from_s3() -> None:
    """
    Download Image Requirements.xlsx from S3 and populate RESIZE_SPECS / CHANNEL_TO_SPEC.

    S3 path: s3://{S3_BUCKET}/Image Requirements.xlsx

    Skipped when S3_BUCKET resolves to an empty string (local unit tests).

    Raises
    ------
    RuntimeError
        If the file cannot be downloaded or parsed.
    """
    global RESIZE_SPECS, CHANNEL_TO_SPEC

    if not S3_BUCKET:
        print(
            "[CONFIG] S3_BUCKET not set - "
            "RESIZE_SPECS and CHANNEL_TO_SPEC will be empty. "
            "Expected during local unit tests only."
        )
        return

    s3_path = f"s3://{S3_BUCKET}/{S3_REQUIREMENTS_KEY}"
    print(f"[CONFIG] Loading requirements from {s3_path} ...")

    try:
        s3_client   = boto3.client("s3")
        response    = s3_client.get_object(
            Bucket=S3_BUCKET,
            Key=S3_REQUIREMENTS_KEY,
        )
        excel_bytes = response["Body"].read()
    except Exception as exc:
        raise RuntimeError(
            f"Failed to download Image Requirements.xlsx from {s3_path}: {exc}. "
            "Check that S3_BUCKET is correct and the Lambda role has "
            "s3:GetObject permission on the bucket."
        ) from exc

    try:
        RESIZE_SPECS, CHANNEL_TO_SPEC = load_config_from_excel(excel_bytes)
    except ValueError as exc:
        raise RuntimeError(
            f"Failed to parse Image Requirements.xlsx: {exc}"
        ) from exc

    print(f"[CONFIG] Loaded {len(RESIZE_SPECS)} spec(s): {sorted(RESIZE_SPECS.keys())}")
    print(f"[CONFIG] Loaded {len(CHANNEL_TO_SPEC)} channel mapping(s).")


# ── Runs once at import time ──────────────────────────────────────────────────
_init_from_s3()