"""
STIBO INBOUND XML GENERATOR -- Implus EAN Source (Sofsole) v2.0
"SUM ORDER" sheet (Sofsole linelist Excel) -> Stibo STEP XML

Purpose:
    Generates Variant-level STEP XML with barcode data.
    Companion to implus/linelist_main_source_sofsole.py (which sends
    Generic-only -- no variants, no barcodes).

v2.0 change (2026-08-26, per explicit spec from the team): simplified
to a direct, PTP-style mapping — replaces v1.0's grouping logic (which
merged multiple SKUs sharing a DESCRIPTION into one generic, keyed off
only the first row's IMPLUS EU ITEM #, to mirror
linelist_main_source_sofsole.py's group_generics()). That grouping
invariant is INTENTIONALLY DROPPED here:

    Principal Style Code -> IMPLUS EU ITEM #   (one row = one generic)
    Size                  -> hardcoded default "000" — no longer read
                              from the SIZE column; every variant uses
                              the same fixed 3-char size code.
    Principal Barcode     -> EAN column (single AT_Barcode entry, same
                              DC_Barcode shape as v1.0)

    Generic (PRD_GenericArticle):
        KEY_InboundArticle  = brand_code + implus_eu
        <Values/>            (intentionally empty — EAN source only)

    Variant (PRD_VariantArticle):
        KEY_InboundVariant  = KEY_InboundArticle + "000"
        <Values/>            (intentionally empty)
        DC_Barcode DataContainer (single entry):
            AT_Barcode           = EAN value
            AT_BarcodeType       = P  (ID-only LOV)
            AT_MainEANIndicator  = Y  (ID-only LOV)

NOTE — this generic key no longer necessarily matches
linelist_main_source_sofsole.py's generic key for rows that script
merges under one generic (multiple IMPLUS EU ITEM # values sharing a
DESCRIPTION land on separate generics here, one per item #). Confirmed
acceptable by the team (2026-08-26) — this file now maps 1 input row
to 1 generic + 1 variant, full stop.

Input sheet: "SUM ORDER" (fallback: active sheet)
Column detection: dynamic, by header name (same signals as
linelist_main_source_sofsole.py):
    IMPLUS EU ITEM #     -> implus_eu  (style/item key -> Principal Style Code)
    EAN                  -> barcode

A row counts as a data row only if it has a non-empty EAN value (banner
/ section-header rows like "INSOLES" and rows with no barcode are
skipped).
"""

from __future__ import annotations

import logging
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

import openpyxl

# ══════════════════════════════════════════════════════════════════════════════
# STIBO / STEP XML CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"

ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z0-9]+)?="[^"]*"')

# Header signal words used to detect the header row in "SUM ORDER" sheet —
# identical set to linelist_main_source_sofsole.py so both modules agree
# on which row is the header.
_HEADER_SIGNALS = {"category", "implus eu", "implus eu item", "upc", "ean", "product description"}

# Size code is fixed for every variant — this file no longer reads a real
# size value (see module docstring v2.0 change).
SIZE_CODE_DEFAULT = "000"

# ══════════════════════════════════════════════════════════════════════════════
# PRODUCT LIMIT — cap the number of generics written to the output XML.
# Each generic has exactly one variant (1 row = 1 generic = 1 variant), so
# this also caps the variant count 1:1. Set to None for no limit.
# ══════════════════════════════════════════════════════════════════════════════
PRODUCT_LIMIT: int | None = None

# ══════════════════════════════════════════════════════════════════════════════
# LOGGING
# ══════════════════════════════════════════════════════════════════════════════
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
log = logging.getLogger("implus.ean_update_source_sofsole")


# ══════════════════════════════════════════════════════════════════════════════
# DATA MODEL
# ══════════════════════════════════════════════════════════════════════════════
@dataclass
class EANRow:
    row_num:    int
    implus_eu:  str
    ean:        str


@dataclass
class VariantInfo:
    size_code:    str
    ean:          str
    variant_key:  str   # KEY_InboundVariant value


@dataclass
class GenericGroup:
    generic_key:  str           # KEY_InboundArticle value
    implus_eu:    str
    variants:     dict[str, VariantInfo] = field(default_factory=dict)
    # key = size_code (always SIZE_CODE_DEFAULT here); last EAN wins on
    # collision (i.e. duplicate IMPLUS EU ITEM # rows)


# ══════════════════════════════════════════════════════════════════════════════
# STRING / BARCODE HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _fmt_barcode(val) -> str:
    """
    Return barcode as a clean string.

    Only run the float→int cleanup (stripping a trailing ".0") on actual
    numeric cell values; string cells are returned as-is so a
    text-formatted barcode with significant leading zeros survives
    untouched.
    """
    if val is None:
        return ""
    if isinstance(val, str):
        s = val.strip()
        return "" if s in ("None", "nan", "NaT") else s
    try:
        f = float(val)
        return str(int(f))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s in ("None", "nan") else s


# ══════════════════════════════════════════════════════════════════════════════
# DIRECTORY HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _get_dirs():
    base     = Path(os.environ.get("LAMBDA_TMP_DIR", str(Path(__file__).parent)))
    in_dir   = base / "input"
    ll_dir   = in_dir / "linelist"
    out_dir  = base / "output"
    xml_dir  = out_dir / "xml"
    log_dir  = out_dir / "logs"
    for d in (in_dir, ll_dir, out_dir, xml_dir, log_dir):
        d.mkdir(parents=True, exist_ok=True)
    return base, in_dir, ll_dir, out_dir, xml_dir, log_dir


def _find_input_files() -> list[Path]:
    _, _, ll_dir, *_ = _get_dirs()
    files: list[Path] = []
    if ll_dir.exists():
        files.extend(
            p for p in ll_dir.iterdir()
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
        )
    return sorted(files)


# ══════════════════════════════════════════════════════════════════════════════
# COLUMN FINDER (identical to linelist_main_source_sofsole.py)
# ══════════════════════════════════════════════════════════════════════════════
def _find_col(header: list[str], *candidates: str) -> Optional[int]:
    """
    Return the 0-based index of the first header cell that matches any
    candidate (case-insensitive, stripped). Returns None if none match.
    """
    normalised = [str(h).strip().lower() for h in header]
    for candidate in candidates:
        target = candidate.strip().lower()
        for i, h in enumerate(normalised):
            if h == target or target in h:
                return i
    return None


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL READER  (header/sheet detection shared with linelist_main_source_sofsole.py)
# ══════════════════════════════════════════════════════════════════════════════
def load_ean_sheet(path: Path) -> list[EANRow]:
    """
    Read EAN / barcode rows from the Sofsole linelist Excel.

    Sheet: "SUM ORDER" (fallback: active sheet).
    Header detection: first row (within the first 30) whose cells match
    at least 2 of _HEADER_SIGNALS — identical to the linelist parser.
    Columns are then resolved dynamically by header name.
    """
    log.info("[EAN] Loading: %s", path.name)
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as exc:
        log.error("[EAN] Cannot open file: %s", exc)
        return []

    ws = wb["SUM ORDER"] if "SUM ORDER" in wb.sheetnames else wb.active
    log.info("[EAN] Using sheet: '%s'", ws.title)

    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        log.warning("[EAN] Sheet is empty")
        return []

    header_row_idx = -1
    for i, row in enumerate(all_rows[:30]):
        if not row:
            continue
        cell_vals = {str(v).strip().lower() for v in row if v is not None}
        matches = sum(1 for sig in _HEADER_SIGNALS if any(sig in cv for cv in cell_vals))
        if matches >= 2:
            header_row_idx = i
            break

    if header_row_idx == -1:
        log.warning("[EAN] Could not detect header row — defaulting to row index 19 (row 20)")
        header_row_idx = 19

    header = [str(h).strip() if h is not None else "" for h in all_rows[header_row_idx]]
    log.info("[EAN] Header row at index %d: %s", header_row_idx, header[:20])

    c_implus_eu = _find_col(header, "IMPLUS EU ITEM #", "IMPLUS EU ITEM", "IMPLUS EU", "SKU", "ITEM #", "ITEM NO")
    c_ean       = _find_col(header, "EAN", "EAN CODE", "EAN/UPC")

    log.info("[EAN] Column indices → implus_eu=%s  ean=%s", c_implus_eu, c_ean)

    def _cell(row, idx) -> str:
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        s = str(row[idx]).strip()
        return "" if s in ("None", "nan") else s

    rows: list[EANRow] = []
    for row_idx, row in enumerate(all_rows[header_row_idx + 1:], start=header_row_idx + 2):
        if not row or all(v is None for v in row):
            continue  # completely blank row

        implus_eu = _cell(row, c_implus_eu)
        if "TOTAL" in implus_eu.upper():
            continue

        ean_val = _fmt_barcode(row[c_ean]) if c_ean is not None and c_ean < len(row) else ""
        if not ean_val:
            # Banner / section-header row (e.g. "INSOLES") or a row with
            # no barcode at all — nothing to write to AT_Barcode.
            continue

        if not implus_eu:
            continue  # no usable join key / Principal Style Code

        rows.append(EANRow(row_num=row_idx, implus_eu=implus_eu, ean=ean_val))

    log.info("[EAN] Parsed %d data rows", len(rows))
    return rows


# ══════════════════════════════════════════════════════════════════════════════
# GROUPING — rows → GenericGroup dict  (1 row = 1 generic = 1 variant)
# ══════════════════════════════════════════════════════════════════════════════
def group_rows(rows: list[EANRow], brand_code: str) -> dict[str, GenericGroup]:
    """
    Group EANRow list into GenericGroup objects.

    Generic key (KEY_InboundArticle) : brand_code + IMPLUS EU ITEM #
    Variant key (KEY_InboundVariant) : generic_key + SIZE_CODE_DEFAULT ("000")

    One input row maps to exactly one generic and one variant — no
    cross-row grouping by description/color (see module docstring v2.0
    change). On generic_key collision (duplicate IMPLUS EU ITEM # rows),
    the last row encountered wins.
    """
    generics: dict[str, GenericGroup] = {}

    for row in rows:
        generic_key = f"{brand_code}{row.implus_eu}"
        variant_key = f"{generic_key}{SIZE_CODE_DEFAULT}"

        if generic_key not in generics:
            generics[generic_key] = GenericGroup(
                generic_key=generic_key,
                implus_eu=row.implus_eu,
            )

        g = generics[generic_key]
        g.variants[SIZE_CODE_DEFAULT] = VariantInfo(
            size_code=SIZE_CODE_DEFAULT,
            ean=row.ean,
            variant_key=variant_key,
        )

    log.info(
        "[Group] %d rows → %d generics, %d variants total",
        len(rows), len(generics), sum(len(g.variants) for g in generics.values()),
    )
    return generics


# ══════════════════════════════════════════════════════════════════════════════
# XML HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _add_barcode_datacontainer(parent: ET.Element, ean: str) -> None:
    """
    Append DC_Barcode DataContainer to a Variant product element.
    Single barcode entry (EAN) — Sofsole's mapping is EAN-only.

        <DataContainers>
          <MultiDataContainer Type="DC_Barcode">
            <DataContainer Analyzer="true">
              <Values>
                <Value AttributeID="AT_Barcode">3700006360159</Value>
                <Value AttributeID="AT_BarcodeType" ID="P"/>
                <Value AttributeID="AT_MainEANIndicator" ID="Y"/>
              </Values>
            </DataContainer>
          </MultiDataContainer>
        </DataContainers>
    """
    if not ean:
        return

    dc_root = ET.SubElement(parent, f"{{{STIBO_NS}}}DataContainers")
    mdc     = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc      = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv     = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = ean

    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


def _build_generic_xml(group: GenericGroup) -> str:
    """
    Build XML string for one Generic article with its single Variant.

        <Product UserTypeID="PRD_GenericArticle">
          <KeyValue KeyID="KEY_InboundArticle">IPL360159</KeyValue>
          <Values/>
          <Product UserTypeID="PRD_VariantArticle">
            <KeyValue KeyID="KEY_InboundVariant">IPL360159000</KeyValue>
            <Values/>
            <DataContainers>...</DataContainers>
          </Product>
        </Product>
    """
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = group.generic_key

    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

    for variant in group.variants.values():
        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant.variant_key

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

        _add_barcode_datacontainer(v_el, variant.ean)

    return ET.tostring(g_el, encoding="unicode")


# ══════════════════════════════════════════════════════════════════════════════
# XML WRITER
# ══════════════════════════════════════════════════════════════════════════════
def build_xml(generics: dict[str, GenericGroup], out_path: Path) -> None:
    export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = 0
    variant_count = 0

    log.info("[EAN] Writing STEP XML → %s", out_path.name)
    with open(out_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n'
        )
        f.write("  <Products>\n")

        for group in generics.values():
            if not group.variants:
                continue
            if PRODUCT_LIMIT is not None and generic_count >= PRODUCT_LIMIT:
                break
            xml_str = _build_generic_xml(group)
            xml_str = _XMLNS_RE.sub("", xml_str)
            f.write(f"    {xml_str}\n")
            generic_count += 1
            variant_count += len(group.variants)

        f.write("  </Products>\n")
        f.write("</STEP-ProductInformation>\n")

    if PRODUCT_LIMIT is not None:
        log.info("[EAN] PRODUCT_LIMIT active: capped at %d generic(s)", PRODUCT_LIMIT)

    log.info(
        "[EAN] Done — %d generics, %d variants written → %s",
        generic_count, variant_count, out_path.name,
    )


# ══════════════════════════════════════════════════════════════════════════════
# ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
def run(args, auditor=None) -> tuple[list[dict], Optional[Path]]:
    """
    Main entry point — called by lambda_function.py.

    args attributes used:
        brand_code   (str)  default "IPL"
        input_file   (str)  optional: explicit path to linelist Excel
    """
    brand_code = getattr(args, "brand_code", "IPL") or "IPL"
    log.info("[EAN-Source] Starting — brand_code=%s", brand_code)

    input_file_arg = getattr(args, "input_file", None)
    if input_file_arg and Path(input_file_arg).exists():
        file_path = Path(input_file_arg)
    else:
        files = _find_input_files()
        if not files:
            log.warning("[EAN-Source] No linelist .xlsx/.xlsm found — nothing to process.")
            return [], None
        file_path = max(files, key=lambda p: p.stat().st_mtime)

    log.info("[EAN-Source] Processing file: %s", file_path.name)

    rows = load_ean_sheet(file_path)
    if not rows:
        log.warning("[EAN-Source] No valid rows parsed from %s", file_path.name)
        return [], None

    generics = group_rows(rows, brand_code)
    if not generics:
        log.warning("[EAN-Source] No valid generics after grouping.")
        return [], None

    _, _, _, _, xml_dir, _ = _get_dirs()
    stem     = file_path.stem
    xml_path = xml_dir / f"{stem}.xml"

    build_xml(generics, xml_path)

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return [], xml_path


# ══════════════════════════════════════════════════════════════════════════════
# CLI (for local testing)
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="Sofsole EAN Source → STEP XML")
    ap.add_argument("--brand-code",  default="IPL", dest="brand_code")
    ap.add_argument("--input-file",  default=None,  dest="input_file",
                    help="Path to Sofsole linelist Excel (optional)")
    run(ap.parse_args())
