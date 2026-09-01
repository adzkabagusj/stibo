"""
STIBO INBOUND XML GENERATOR -- Impulse EAN Update v1.0
"Balega Asia Gen" sheet (same linelist Excel) -> Stibo STEP XML

Purpose:
    Generates Variant-level STEP XML with barcode data.
    Companion to impulse/linelist_main.py (which sends Generic-only).

MAPPED ATTRIBUTES:
    Generic (PRD_GenericArticle):
        KEY_InboundArticle  = brand_code + implus_us
        <Values/>            (intentionally empty — EAN update only)

    Variant (PRD_VariantArticle):
        KEY_InboundVariant  = KEY_InboundArticle + size_code (3-char)
        <Values/>            (intentionally empty)
        DC_Barcode DataContainer:
            AT_Barcode           = UPC / barcode value (col 2)
            AT_BarcodeType       = P  (ID-only LOV)
            AT_MainEANIndicator  = Y  (ID-only LOV)

KEY_InboundVariant formula:
    KEY_InboundArticle + _size_code(size)
    e.g. brand_code="IPL", implus_us="F010000", size="M"
    → "IPLF010000" + "00M" = "IPLF01000000M"

Input sheet: "Balega Asia Gen" (or active sheet if not found)
Column layout (0-indexed, same as linelist_main.py):
    col 0  = Color Code   -> color_code
    col 1  = Implus US    -> implus_us  (style/item key)
    col 2  = UPC          -> barcode
    col 4  = Style Desc   -> (not used in EAN update XML)
    col 5  = Color        -> color (grouping only)
    col 6  = Size         -> size -> size_code for KEY_InboundVariant
"""

from __future__ import annotations

import logging
import os
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

import openpyxl

# ══════════════════════════════════════════════════════════════════════════════
# STIBO / STEP XML CONSTANTS
# ══════════════════════════════════════════════════════════════════════════════
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"

ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z0-9]+)?="[^"]*"')

# ══════════════════════════════════════════════════════════════════════════════
# LOGGING
# ══════════════════════════════════════════════════════════════════════════════
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
log = logging.getLogger("impulse.ean_update_main")


# ══════════════════════════════════════════════════════════════════════════════
# DATA MODEL
# ══════════════════════════════════════════════════════════════════════════════
@dataclass
class EANRow:
    row_num:    int
    implus_us:  str = ""
    color_code: str = ""
    color:      str = ""
    size:       str = ""
    upc:        str = ""


@dataclass
class VariantInfo:
    size_raw:     str
    size_code:    str
    upc:          str
    variant_key:  str   # KEY_InboundVariant value


@dataclass
class GenericGroup:
    generic_key:  str           # KEY_InboundArticle value
    implus_us:    str
    variants:     dict[str, VariantInfo] = field(default_factory=dict)
    # key = size_code (deduplication); last UPC wins on collision


# ══════════════════════════════════════════════════════════════════════════════
# SIZE CODE FORMATTER
# ══════════════════════════════════════════════════════════════════════════════
_SIZE_OVERRIDES: dict[str, str] = {
    "ONE SIZE": "ONS",
    "ONE-SIZE": "ONS",
    "OS":       "ONS",
    "ONESIZE":  "ONS",
    "FREE":     "ONS",
    "FREE SIZE": "ONS",
    "NS":       "NSZ",
    "NO SIZE":  "NSZ",
}

# Word-size aliases — hardcoded per user direction (2026-08-17): spelled-out
# sizes on Implus EAN files (Small/Medium/Large, etc.) collapse to their
# standard letter abbreviation BEFORE the normal padding rules run, e.g.
# "Small" -> "S" -> "00S" (same padding a literal "S" already gets),
# "X-Large" -> "XL" -> "0XL", "XX-Large" -> "XXL" (already 3 chars).
_SIZE_WORD_ALIASES: dict[str, str] = {
    "SMALL":              "S",
    "MEDIUM":             "M",
    "LARGE":              "L",
    "X-SMALL":            "XS",
    "XSMALL":             "XS",
    "EXTRA SMALL":        "XS",
    "EXTRA-SMALL":        "XS",
    "X-LARGE":            "XL",
    "XLARGE":             "XL",
    "EXTRA LARGE":        "XL",
    "EXTRA-LARGE":        "XL",
    "XX-LARGE":           "XXL",
    "XXLARGE":            "XXL",
    "2X LARGE":           "XXL",
    "2X-LARGE":           "XXL",
    "2XLARGE":            "XXL",
    "DOUBLE EXTRA LARGE": "XXL",
}


def _size_code(size_val: str) -> str:
    """
    Convert a raw size string to a 3-char MAA size code for use in
    KEY_InboundVariant.

    Priority:
      0. Word-size alias (Small/Medium/Large/X-Small/X-Large/XX-Large, etc.)
         → standard letter abbreviation, THEN the normal rules below
         (e.g. "Small" → "S" → "00S"; "X-Large" → "XL" → "0XL")
      1. Override table (ONE SIZE → ONS, etc.)
      2. Pure integer → zero-padded to 3 digits (e.g. "9" → "009")
      3. Half-size decimal (e.g. "9.5") → "09H"
      4. Other decimal → leading integer zero-padded (e.g. "12.0" → "012")
      5. 1-char alpha → "00X"  (e.g. "M" → "00M")
      6. 2-char alpha → "0XY"  (e.g. "XL" → "0XL")
      7. 3+ char alpha/alnum → first 3 chars uppercased (e.g. "XXL" → "XXL")
    """
    if not size_val:
        return "MSC"
    s = str(size_val).strip().upper()
    if not s:
        return "MSC"

    # 0. Word-size alias — spelled-out sizes collapse to their letter form
    # before anything else runs.
    s = _SIZE_WORD_ALIASES.get(s, s)

    # 1. Override table
    if s in _SIZE_OVERRIDES:
        return _SIZE_OVERRIDES[s]

    # 2. Pure integer
    if re.match(r"^\d+$", s):
        try:
            return str(int(s)).zfill(3)[:3]
        except ValueError:
            pass

    # 3. Half-size (e.g. 9.5)
    m = re.match(r"^(\d+)\.5$", s)
    if m:
        return str(int(m.group(1))).zfill(2)[:2] + "H"

    # 4. Other decimal (e.g. 9.0, 12.33)
    m2 = re.match(r"^(\d+)\.\d+$", s)
    if m2:
        return str(int(m2.group(1))).zfill(3)[:3]

    # 5-7. Alpha / alnum
    alnum = re.sub(r"[^A-Z0-9]", "", s)
    if not alnum:
        return "MSC"
    if len(alnum) == 1:
        return "00" + alnum
    if len(alnum) == 2:
        return "0" + alnum
    return alnum[:3]


# ══════════════════════════════════════════════════════════════════════════════
# STRING HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _s(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "NaT") else s


def _fmt_barcode(val) -> str:
    """Return barcode as a clean integer string (strips .0 from floats)."""
    if val is None:
        return ""
    try:
        f = float(str(val).strip())
        return str(int(f))
    except (ValueError, TypeError):
        s = str(val).strip()
        return "" if s in ("None", "nan") else s


# ══════════════════════════════════════════════════════════════════════════════
# DIRECTORY HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _get_dirs():
    base     = Path(os.environ.get("LAMBDA_TMP_DIR", str(Path(__file__).parent)))
    in_dir   = base / "input"
    ll_dir   = in_dir / "linelist"
    out_dir  = base / "output"
    xml_dir  = out_dir / "xml"
    log_dir  = out_dir / "logs"
    for d in (in_dir, ll_dir, out_dir, xml_dir, log_dir):
        d.mkdir(parents=True, exist_ok=True)
    return base, in_dir, ll_dir, out_dir, xml_dir, log_dir


def _find_input_files() -> list[Path]:
    _, _, ll_dir, *_ = _get_dirs()
    files: list[Path] = []
    if ll_dir.exists():
        files.extend(
            p for p in ll_dir.iterdir()
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm")
        )
    return sorted(files)


# ══════════════════════════════════════════════════════════════════════════════
# EXCEL READER  (same sheet as linelist — "Balega Asia Gen")
# ══════════════════════════════════════════════════════════════════════════════
def load_ean_sheet(path: Path) -> list[EANRow]:
    """
    Read EAN / barcode rows from the Impulse linelist Excel.

    Sheet: "Balega Asia Gen" (fallback: active sheet).
    Header detection: looks for a row where col 0 = "Color Code"
                      or col 1 = "Implus US".
    Data columns (0-indexed):
        0  Color Code  → color_code
        1  Implus US   → implus_us
        2  UPC         → upc (barcode)
        5  Color       → color
        6  Size        → size
    """
    log.info("[EAN] Loading: %s", path.name)
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as exc:
        log.error("[EAN] Cannot open file: %s", exc)
        return []

    ws = wb["Balega Asia Gen"] if "Balega Asia Gen" in wb.sheetnames else wb.active
    log.info("[EAN] Using sheet: '%s'", ws.title)

    def _c(row, idx) -> str:
        if idx >= len(row) or row[idx] is None:
            return ""
        s = str(row[idx]).strip()
        return "" if s in ("None", "nan") else s

    # Detect header row
    header_row_idx = -1
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        c0 = _c(row, 0).lower()
        c1 = _c(row, 1).lower()
        if c0 == "color code" or c1 == "implus us":
            header_row_idx = row_idx
            break

    if header_row_idx == -1:
        log.warning("[EAN] Header row not found — defaulting to row 20")
        header_row_idx = 20

    rows: list[EANRow] = []
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row_idx + 1, values_only=True),
        start=header_row_idx + 1,
    ):
        c0 = _c(row, 0)
        c1 = _c(row, 1)
        c2 = _c(row, 2) if len(row) > 2 else ""
        c4 = _c(row, 4) if len(row) > 4 else ""

        # Section banner rows: only col 0/1 have text, col 2 and col 4 are empty
        if (c0 or c1) and not c2 and not c4:
            continue
        if not c0 and not c1:
            continue

        implus_us = _c(row, 1)
        if not implus_us:
            continue

        color_code = _c(row, 0)
        upc        = _fmt_barcode(row[2] if len(row) > 2 else None)
        color      = _c(row, 5) if len(row) > 5 else ""
        size       = _c(row, 6) if len(row) > 6 else ""

        rows.append(EANRow(
            row_num=row_idx,
            implus_us=implus_us,
            color_code=color_code,
            color=color,
            size=size,
            upc=upc,
        ))

    wb.close()
    log.info("[EAN] Parsed %d data rows", len(rows))
    return rows


# ══════════════════════════════════════════════════════════════════════════════
# GROUPING — rows → GenericGroup dict
# ══════════════════════════════════════════════════════════════════════════════
def group_rows(rows: list[EANRow], brand_code: str) -> dict[str, GenericGroup]:
    """
    Group EANRow list into GenericGroup objects.

    Generic key: brand_code + implus_us
    Variant key: generic_key + size_code (3-char)

    On size_code collision within a generic the last UPC encountered wins
    (same approach as dr_marten inline_ean_source.py).
    """
    generics: dict[str, GenericGroup] = {}

    for row in rows:
        generic_key = f"{brand_code}{row.implus_us}"
        sc          = _size_code(row.size)
        variant_key = f"{generic_key}{sc}"

        if generic_key not in generics:
            generics[generic_key] = GenericGroup(
                generic_key=generic_key,
                implus_us=row.implus_us,
            )

        g = generics[generic_key]
        g.variants[sc] = VariantInfo(
            size_raw=row.size,
            size_code=sc,
            upc=row.upc,
            variant_key=variant_key,
        )

    log.info(
        "[Group] %d rows → %d generics, %d variants total",
        len(rows),
        len(generics),
        sum(len(g.variants) for g in generics.values()),
    )
    return generics


# ══════════════════════════════════════════════════════════════════════════════
# XML HELPERS
# ══════════════════════════════════════════════════════════════════════════════
def _add_barcode_datacontainer(parent: ET.Element, ean: str) -> None:
    """
    Append DC_Barcode DataContainer to a Variant product element.

        <DataContainers>
          <MultiDataContainer Type="DC_Barcode">
            <DataContainer Analyzer="true">
              <Values>
                <Value AttributeID="AT_Barcode">190665653885</Value>
                <Value AttributeID="AT_BarcodeType" ID="P"/>
                <Value AttributeID="AT_MainEANIndicator" ID="Y"/>
              </Values>
            </DataContainer>
          </MultiDataContainer>
        </DataContainers>
    """
    if not ean:
        return

    dc_root = ET.SubElement(parent, f"{{{STIBO_NS}}}DataContainers")
    mdc     = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc      = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv     = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = ean

    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


def _build_generic_xml(group: GenericGroup) -> str:
    """
    Build XML string for one Generic article with all its Variants.

        <Product UserTypeID="PRD_GenericArticle">
          <KeyValue KeyID="KEY_InboundArticle">IPLF010000</KeyValue>
          <Values/>
          <Product UserTypeID="PRD_VariantArticle">
            <KeyValue KeyID="KEY_InboundVariant">IPLF01000000M</KeyValue>
            <Values/>
            <DataContainers>...</DataContainers>
          </Product>
          ...
        </Product>
    """
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    kv.set("KeyID", "KEY_InboundArticle")
    kv.text = group.generic_key

    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

    for variant in group.variants.values():
        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant.variant_key

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")  # intentionally empty

        _add_barcode_datacontainer(v_el, variant.upc)

    return ET.tostring(g_el, encoding="unicode")


# ══════════════════════════════════════════════════════════════════════════════
# XML WRITER
# ══════════════════════════════════════════════════════════════════════════════
def build_xml(generics: dict[str, GenericGroup], out_path: Path) -> None:
    export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    generic_count = 0
    variant_count = 0

    # ══════════════════════════════════════════════════════════════════════
    # TEST LIMITER — Set TEST_MODE = False for production
    # ══════════════════════════════════════════════════════════════════════
    TEST_MODE = False
    PRODUCT_LIMIT = 2
    VARIANT_LIMIT = 3

    log.info("[EAN] Writing STEP XML → %s", out_path.name)
    with open(out_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n'
        )
        f.write("  <Products>\n")

        for group in generics.values():
            if not group.variants:
                continue
            if TEST_MODE:
                if generic_count >= PRODUCT_LIMIT:
                    break
                if variant_count + len(group.variants) > VARIANT_LIMIT:
                    continue  # would overflow the variant budget — try the next group
            xml_str = _build_generic_xml(group)
            xml_str = _XMLNS_RE.sub("", xml_str)
            f.write(f"    {xml_str}\n")
            generic_count += 1
            variant_count += len(group.variants)

        f.write("  </Products>\n")
        f.write("</STEP-ProductInformation>\n")

    if TEST_MODE:
        log.info("[XML] TEST MODE: limited to %d product(s), %d variant(s)", PRODUCT_LIMIT, VARIANT_LIMIT)

    log.info(
        "[EAN] Done — %d generics, %d variants written → %s",
        generic_count, variant_count, out_path.name,
    )


# ══════════════════════════════════════════════════════════════════════════════
# ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
def run(args, auditor=None) -> tuple[list[dict], Optional[Path]]:
    """
    Main entry point — called by lambda_function.py.

    args attributes used:
        brand_code   (str)  default "IPL"
        input_file   (str)  optional: explicit path to linelist Excel
    """
    brand_code = getattr(args, "brand_code", "IPL") or "IPL"
    log.info("[EAN-Update] Starting — brand_code=%s", brand_code)

    input_file_arg = getattr(args, "input_file", None)
    if input_file_arg and Path(input_file_arg).exists():
        file_path = Path(input_file_arg)
    else:
        files = _find_input_files()
        if not files:
            log.warning("[EAN-Update] No linelist .xlsx/.xlsm found — nothing to process.")
            return [], None
        file_path = max(files, key=lambda p: p.stat().st_mtime)

    log.info("[EAN-Update] Processing file: %s", file_path.name)

    rows = load_ean_sheet(file_path)
    if not rows:
        log.warning("[EAN-Update] No valid rows parsed from %s", file_path.name)
        return [], None

    generics = group_rows(rows, brand_code)
    if not generics:
        log.warning("[EAN-Update] No valid generics after grouping.")
        return [], None

    _, _, _, _, xml_dir, _ = _get_dirs()
    stem     = file_path.stem
    xml_path = xml_dir / f"{stem}_EAN_Update.xml"

    build_xml(generics, xml_path)

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return [], xml_path


# ══════════════════════════════════════════════════════════════════════════════
# CLI (for local testing)
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="Impulse EAN Update → STEP XML")
    ap.add_argument("--brand-code",  default="IPL", dest="brand_code")
    ap.add_argument("--input-file",  default=None,  dest="input_file",
                    help="Path to Impulse linelist Excel (optional)")
    run(ap.parse_args())
