"""
╔══════════════════════════════════════════════════════════════════╗
║ STIBO INBOUND XML GENERATOR — New Era Delivery Schedule EAN v1 ║
║ Delivery Schedule EAN → Stibo STEP XML (Generic + Variants)    ║
╚══════════════════════════════════════════════════════════════════╝

Input sheet:
    Sheet1

Column mapping:
    Material Number  -> Principal Style Code (last 8 digits)
    Grid Size        -> Size display value (lookup in MDD Size Code LOV)
    EAN/UPC          -> Barcode

Rules:
  • Principal style code = last 8 digits from "Material Number".
  • Size LOV ID must be resolved from MDD "Size Code LOV".
  • NO fallback formula for size. If LOV ID not found, skip that variant.
  • One Generic per principal style code.
  • One Variant per size LOV ID.
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

# ======================================================================
# PATHS
# ======================================================================
BASE_DIR = Path(os.environ.get("LAMBDA_TMP_DIR", Path(__file__).resolve().parent))

EAN_DIR     = BASE_DIR / "input" / "delivery_schedule_ean"
MDD_DIR     = BASE_DIR / "input" / "mdd"
XML_OUT_DIR = BASE_DIR / "output" / "xml"
LOG_DIR     = BASE_DIR / "output" / "logs"

for _d in (EAN_DIR, MDD_DIR, XML_OUT_DIR, LOG_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# ======================================================================
# BRAND CONSTANTS
# ======================================================================
BRAND_NAME = "New Era"
BRAND_CODE = "NRA"
INPUT_SHEET_NAME = "Sheet1"

# ======================================================================
# XML NAMESPACE
# ======================================================================
STIBO_NS     = "http://www.stibosystems.com/step"
STIBO_XSI    = "http://www.w3.org/2001/XMLSchema-instance"
STIBO_SCHEMA = "http://www.stibosystems.com/step PIM.xsd"
ET.register_namespace("", STIBO_NS)
_XMLNS_RE = re.compile(r'\s+xmlns(?::[a-z]+)?="[^"]*"')


# ======================================================================
# HELPERS
# ======================================================================

def _s(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("", "None", "nan", "NaT", "#N/A", "#VALUE!") else s


def _norm_header(s: str) -> str:
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _list_excel_files(folder: Path) -> list[Path]:
    """Case-insensitive Excel file listing for Linux Lambda runtime."""
    if not folder.exists():
        return []
    return sorted(
        [
            p for p in folder.iterdir()
            if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".xls")
        ],
        key=lambda p: p.name.lower(),
    )


def _fmt_barcode(v) -> str:
    if v is None:
        return ""
    try:
        return str(int(float(str(v).strip())))
    except (ValueError, TypeError):
        return _s(v)


def _extract_principal_style_code(material_number) -> str:
    """Take last 8 digits from Material Number. Returns '' if unavailable."""
    raw = _s(material_number)
    if not raw:
        return ""
    digits = "".join(re.findall(r"\d", raw))
    if len(digits) < 8:
        return ""
    return digits[-8:]


def _transform_grid_size(size_val: str) -> str:
    """Normalize M/W format for lookup (e.g. M7W9 -> M7/W9)."""
    s = _s(size_val).upper()
    if not s:
        return ""
    m = re.match(r"^(M\d+)(W\d+)$", s)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    return s


def _resolve_size_lov_id(size_val: str, size_lov: dict[str, str] | None) -> str:
    """Resolve size LOV ID strictly from MDD. No fallback formula."""
    if not size_lov:
        return ""

    s = _transform_grid_size(size_val)
    if not s:
        return ""

    candidates = [
        s,
        s.replace(" ", ""),
        s.replace("/", ""),
        (s.lstrip("0") or "0"),
    ]

    for c in candidates:
        if c in size_lov:
            return size_lov[c]
    return ""


# ======================================================================
# LOADERS
# ======================================================================

class MDDLoader:
    """Loads Size Code LOV (Col G description -> Col F LOV ID)."""

    def __init__(self, path: Path):
        self.path = path
        self.lovs: dict[str, dict[str, str]] = {}
        self._load()

    def _load(self) -> None:
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        self._load_size_lov(wb)
        wb.close()

    def _load_size_lov(self, wb) -> None:
        if "Size Code LOV" not in wb.sheetnames:
            log.warning("[MDD] Size Code LOV sheet not found")
            return

        rows = list(wb["Size Code LOV"].iter_rows(values_only=True))
        size_map: dict[str, str] = {}
        for row in rows[1:]:
            if not row or len(row) < 7:
                continue
            lov_id = row[5]  # Col F
            desc = row[6]    # Col G
            if lov_id is None or desc is None:
                continue
            lid = _s(lov_id)
            key = _s(desc).upper().rstrip(".")
            if key and lid and key not in size_map:
                size_map[key] = lid

        self.lovs["Size Code"] = size_map


class NewEraDeliveryScheduleEANLoader:
    TARGET_SHEET = INPUT_SHEET_NAME

    def __init__(self, path: Path):
        self.path = path
        self.df = pd.DataFrame()
        self.sheet = ""
        self.headers: list[str] = []
        self._load()

    def _load(self) -> None:
        log.info("[NewEra-DeliveryScheduleEAN] Loading: %s", self.path.name)
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)

        target = next((s for s in wb.sheetnames if s.strip().upper() == self.TARGET_SHEET.upper()), None)
        if not target:
            wb.close()
            raise ValueError(
                f"[NewEra-DeliveryScheduleEAN] Sheet '{self.TARGET_SHEET}' not found. "
                f"Available: {wb.sheetnames}"
            )

        ws = wb[target]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            self.df = pd.DataFrame()
            self.sheet = target
            self.headers = []
            return

        header_signals = {"material number", "grid size", "ean/upc", "ean", "upc"}
        hdr_idx = 0
        for i in range(min(20, len(rows))):
            r = rows[i]
            if not r:
                continue
            rv = {_norm_header(v) for v in r if v is not None}
            if rv & header_signals:
                hdr_idx = i
                break

        header = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[hdr_idx])]

        seen: dict[str, int] = {}
        for i, col in enumerate(header):
            if col in seen:
                seen[col] += 1
                header[i] = f"{col}_{seen[col]}"
            else:
                seen[col] = 0

        df = pd.DataFrame(rows[hdr_idx + 1:], columns=header)

        material_col = self._find_column(header, ["Material Number", "Material", "Material No", "Material No."])
        if material_col:
            df = df[df[material_col].notna() & (~df[material_col].astype(str).str.strip().isin(["", "None", "nan"]))]

        self.df = df.reset_index(drop=True)
        self.sheet = target
        self.headers = header

    def _find_column(self, headers: list[str], candidates: list[str]) -> str | None:
        for c in candidates:
            for h in headers:
                if _norm_header(h) == _norm_header(c):
                    return h
        return None

    def find_column(self, *candidates: str) -> str | None:
        return self._find_column(self.headers, list(candidates))


# ======================================================================
# MAPPER
# ======================================================================

def group_rows(raw_rows: list[dict], brand_code: str,
               loader: NewEraDeliveryScheduleEANLoader,
               mdd: MDDLoader | None = None) -> dict[str, dict]:
    """
    One Generic per principal style code.
    One Variant per Size LOV ID (strict MDD lookup; missing LOV -> skip).
    """
    result: dict[str, dict] = {}

    material_col = loader.find_column("Material Number", "Material", "Material No", "Material No.")
    grid_size_col = loader.find_column("Grid Size", "Size")
    ean_col = loader.find_column("EAN/UPC", "EAN", "UPC", "UPC Code", "Barcode", "GTIN")

    if not material_col:
        raise ValueError("[NewEra-DeliveryScheduleEAN] Cannot find 'Material Number' column")
    if not grid_size_col:
        raise ValueError("[NewEra-DeliveryScheduleEAN] Cannot find 'Grid Size' column")
    if not ean_col:
        log.warning("[NewEra-DeliveryScheduleEAN] 'EAN/UPC' column not found — barcodes will be empty")

    size_lov = (mdd.lovs.get("Size Code") if mdd else None) or {}

    for row in raw_rows:
        style_code = _extract_principal_style_code(row.get(material_col))
        if not style_code:
            continue

        size_raw = _s(row.get(grid_size_col))
        if not size_raw:
            continue

        size_lov_id = _resolve_size_lov_id(size_raw, size_lov)
        if not size_lov_id:
            # Strict rule requested: no fallback -> skip variant
            continue

        barcode = _fmt_barcode(row.get(ean_col)) if ean_col else ""

        generic_code = f"{brand_code}{style_code}"
        if generic_code not in result:
            result[generic_code] = {
                "style_code": style_code,
                "generic_code": generic_code,
                "variants": {},
            }

        variant_key = size_lov_id
        if variant_key not in result[generic_code]["variants"]:
            result[generic_code]["variants"][variant_key] = {
                "size_raw": size_raw,
                "size_lov_id": size_lov_id,
                "barcode": barcode,
                "variant_code": f"{generic_code}{size_lov_id}",
            }

    return result


def _add_barcode_datacontainer(parent_el: ET.Element, barcode: str) -> None:
    if not barcode:
        return

    dc_root = ET.SubElement(parent_el, f"{{{STIBO_NS}}}DataContainers")
    mdc = ET.SubElement(dc_root, f"{{{STIBO_NS}}}MultiDataContainer")
    mdc.set("Type", "DC_Barcode")
    dc = ET.SubElement(mdc, f"{{{STIBO_NS}}}DataContainer")
    dc.set("Analyzer", "true")
    dcv = ET.SubElement(dc, f"{{{STIBO_NS}}}Values")

    v1 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v1.set("AttributeID", "AT_Barcode")
    v1.text = barcode

    v2 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v2.set("AttributeID", "AT_BarcodeType")
    v2.set("ID", "P")

    v3 = ET.SubElement(dcv, f"{{{STIBO_NS}}}Value")
    v3.set("AttributeID", "AT_MainEANIndicator")
    v3.set("ID", "Y")


# ======================================================================
# XML BUILDER
# ======================================================================

def build_product_xml(generic: dict) -> str:
    g_el = ET.Element(f"{{{STIBO_NS}}}Product")
    g_el.set("UserTypeID", "PRD_GenericArticle")

    g_kv = ET.SubElement(g_el, f"{{{STIBO_NS}}}KeyValue")
    g_kv.set("KeyID", "KEY_InboundArticle")
    g_kv.text = generic["generic_code"]

    # Keep generic <Values> present but empty, matching Crocs shipment flow.
    ET.SubElement(g_el, f"{{{STIBO_NS}}}Values")

    for size_lov_id in sorted(generic["variants"].keys()):
        variant = generic["variants"][size_lov_id]

        v_el = ET.SubElement(g_el, f"{{{STIBO_NS}}}Product")
        v_el.set("UserTypeID", "PRD_VariantArticle")

        v_kv = ET.SubElement(v_el, f"{{{STIBO_NS}}}KeyValue")
        v_kv.set("KeyID", "KEY_InboundVariant")
        v_kv.text = variant["variant_code"]

        ET.SubElement(v_el, f"{{{STIBO_NS}}}Values")

        if variant.get("barcode"):
            _add_barcode_datacontainer(v_el, variant["barcode"])

    return ET.tostring(g_el, encoding="unicode")


# ======================================================================
# ORCHESTRATOR
# ======================================================================

def run(args, auditor=None):
    brand = getattr(args, "brand", BRAND_NAME) or BRAND_NAME
    brand_code = getattr(args, "brand_code", BRAND_CODE) or BRAND_CODE

    mdd = None
    mdd_file = next(iter(_list_excel_files(MDD_DIR)), None)
    if mdd_file:
        mdd = MDDLoader(mdd_file)
    else:
        log.warning("[NewEra-DeliveryScheduleEAN] No MDD file found in %s", MDD_DIR)

    ean_file = next(iter(_list_excel_files(EAN_DIR)), None)
    if not ean_file:
        raise FileNotFoundError(f"No New Era Delivery Schedule EAN file found in {EAN_DIR}")

    loader = NewEraDeliveryScheduleEANLoader(ean_file)
    if loader.df.empty:
        log.warning("[NewEra-DeliveryScheduleEAN] No data rows found")
        return

    raw_rows = loader.df.to_dict("records")
    generics = group_rows(raw_rows, brand_code, loader, mdd=mdd)
    if not generics:
        log.warning("[NewEra-DeliveryScheduleEAN] No valid generics/variants produced")
        return

    xml_filename = f"{ean_file.stem}.xml"
    xml_path = XML_OUT_DIR / xml_filename
    export_time = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

    generic_count = 0
    variant_count = 0

    with open(xml_path, "w", encoding="utf-8", buffering=1 << 20) as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n')
        f.write(
            f'<STEP-ProductInformation '
            f'xmlns="{STIBO_NS}" '
            f'xmlns:xsi="{STIBO_XSI}" '
            f'xsi:schemaLocation="{STIBO_SCHEMA}" '
            f'ExportTime="{export_time}" '
            f'ExportContext="Context1" '
            f'WorkspaceID="Main" '
            f'UseContextLocale="false">\n\n'
        )
        f.write("  <Products>\n")

        for generic in generics.values():
            px = build_product_xml(generic)
            if px:
                f.write(f"    {_XMLNS_RE.sub('', px)}\n")
                generic_count += 1
                variant_count += len(generic["variants"])

        f.write("  </Products>\n</STEP-ProductInformation>\n")

    file_kb = xml_path.stat().st_size // 1024
    print(
        f"=== NEW ERA DELIVERY SCHEDULE EAN SUMMARY ===\n"
        f"  Input rows : {len(raw_rows)}\n"
        f"  Generics   : {generic_count}\n"
        f"  Variants   : {variant_count}\n"
        f"  XML size   : {file_kb} KB",
        flush=True,
    )

    if auditor:
        auditor.set_xml_uploads([str(xml_path)])

    return xml_path


# ======================================================================
# FILENAME METADATA PARSER
# ======================================================================

def _parse_metadata_from_filename(dirs: dict, triggered_file_type: str = None) -> dict:
    candidate_dirs = []
    if triggered_file_type and triggered_file_type in dirs:
        candidate_dirs.append(dirs[triggered_file_type])
    candidate_dirs.append(dirs.get("delivery_schedule_ean"))

    seen: set[str] = set()
    candidate_dirs = [
        d for d in candidate_dirs
        if d is not None and str(d) not in seen and not seen.add(str(d))
    ]

    for folder in candidate_dirs:
        if not folder or not folder.exists():
            continue
        for xlsx_file in _list_excel_files(folder):
            stem = xlsx_file.stem
            parts = re.split(r"\s*-\s*", stem)
            if len(parts) < 6:
                continue

            comp_code = parts[0].strip()
            sbu = parts[1].strip()
            brand = parts[2].strip().title() if len(parts) > 2 else BRAND_NAME

            season = None
            season_idx = None
            for i, p in enumerate(parts):
                if re.match(r"^[A-Z]{2}\d{2,4}$", p.strip(), re.IGNORECASE):
                    season = p.strip().upper()
                    season_idx = i
                    break

            if not season or season_idx is None:
                continue

            trailing = parts[season_idx + 1:]

            seq = 1
            for p in reversed(trailing):
                if p.strip().isdigit():
                    seq = int(p.strip())
                    break

            country_code = ""
            for p in trailing:
                p = p.strip()
                if re.match(r"^[A-Z]{2,3}$", p, re.IGNORECASE) and not p.isdigit():
                    country_code = p.upper()
                    break

            return {
                "comp_code": comp_code or "0888",
                "sbu": sbu or "SP",
                "brand": brand,
                "brand_code": BRAND_CODE,
                "season": season,
                "seq": seq,
                "country_code": country_code,
            }

    return {
        "comp_code": "0888",
        "sbu": "SP",
        "brand": BRAND_NAME,
        "brand_code": BRAND_CODE,
        "season": "SS27",
        "seq": 1,
        "country_code": "",
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="New Era Delivery Schedule EAN -> Stibo XML")
    parser.add_argument("--brand", default=BRAND_NAME)
    parser.add_argument("--brand_code", default=BRAND_CODE)
    parser.add_argument("--comp_code", default="0888")
    parser.add_argument("--sbu", default="SP")
    parser.add_argument("--season", default="SS27")
    parser.add_argument("--country_code", default="")
    cli_args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-8s  %(message)s",
        handlers=[logging.StreamHandler(sys.stdout)],
    )
    run(cli_args)
